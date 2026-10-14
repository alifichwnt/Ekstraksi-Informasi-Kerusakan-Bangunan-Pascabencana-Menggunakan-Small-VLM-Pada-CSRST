#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
11_buat_formulir_sesi_a.py
==========================
Formulir evaluator mahasiswa SESI A (tujuan 2, evaluasi otomatis).

Evaluator menuliskan jawaban deskriptif (satu kalimat) tentang kondisi bangunan
sasaran pada 40 bangunan sampel kecil. Kalimat-kalimat ini menjadi DESKRIPSI
REFERENSI (multiple references, 5 evaluator) untuk metrik BLEU, METEOR, dan
ROUGE-1/2/L terhadap deskripsi keluaran model.

Kesetaraan dengan model:
  * citra sama (chip gabung: kiri sebelum, kanan sesudah bencana);
  * definisi 4 fitur terpilih diambil dari konfigurasi_fitur.json yang sama dengan
    prompt, sebagai pemandu isi deskripsi;
  * panjang jawaban disamakan dengan instruksi model: SATU kalimat.

Penilaian buta: evaluator tidak diberi label xBD dan tidak melihat jawaban model.
Sesi A wajib kembali SEBELUM sesi B (penilaian keluaran model) dibagikan.
Urutan bangunan diacak (benih 42) dengan kode E01-E40; kunci kode -> uid disimpan
di kunci_formulir_sesi_a.csv, hanya untuk peneliti.

Penggunaan (dari folder KOMPONEN_4):
    pip install pandas openpyxl pillow
    python 11_buat_formulir_sesi_a.py
Keluaran: KOMPONEN_6\\Formulir\\Formulir Sesi A - EVAL01.xlsx ... EVAL05.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from pathlib import Path

try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from PIL import Image
except ImportError:  # pragma: no cover
    sys.exit("Butuh pandas, openpyxl, pillow. Jalankan:  pip install pandas openpyxl pillow")

AKAR = Path(os.environ.get("SKRIPSI_AKAR",
             r"F:\KULIAH STIS\TINGKAT IV - SD\SKRIPSWEET\MANUVER SKRIPSI AFTER SENSUS"))
SAMPEL_KECIL = AKAR / "KOMPONEN_3" / "Output Sampel" / "sampel_kecil.csv"
GABUNG = AKAR / "KOMPONEN_3" / "Output Chip" / "gabung"
KONFIG = AKAR / "KOMPONEN_4" / "konfigurasi_fitur.json"
OUT = AKAR / "KOMPONEN_6" / "Formulir"

BENIH = 42
LEBAR_GAMBAR = 520
TIPIS = Side(style="thin", color="999999")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
ISI = PatternFill("solid", fgColor="FFF7D6")
KEPALA = PatternFill("solid", fgColor="DCE6F1")
ATAS = Alignment(wrap_text=True, vertical="top")

# setara dengan aturan deskripsi pada prompt_fitur_v4 (04_susun_prompt.py, ATURAN_FITUR #5)
INSTRUKSI = ("Tulis SATU kalimat bahasa Indonesia dengan kata-kata Anda sendiri tentang kondisi "
             "bangunan sasaran pada citra sesudah bencana, dibandingkan sebelum bencana. "
             "Bila tampak, sebutkan hal-hal pada daftar pemandu di bawah. Jangan menyalin kalimat "
             "definisi.")


def main() -> None:
    a = argparse.ArgumentParser()
    a.add_argument("--out", type=Path, default=OUT)
    a.add_argument("--jumlah", type=int, default=5, help="jumlah evaluator")
    arg = a.parse_args()

    k = json.loads(KONFIG.read_text(encoding="utf-8"))
    peta = {c["kunci"]: c for c in k["kandidat"]}
    fitur = [peta[t] for t in k["seleksi"]["terpilih"]]

    s = pd.read_csv(SAMPEL_KECIL).sample(frac=1, random_state=BENIH).reset_index(drop=True)
    s["kode"] = [f"E{i:02d}" for i in range(1, len(s) + 1)]
    s["f_gabung"] = s["chip_pra"].str.replace("_pra", "_gabung", regex=False).map(lambda n: GABUNG / n)
    hilang = s[~s["f_gabung"].map(Path.exists)]
    if len(hilang):
        sys.exit(f"{len(hilang)} chip gabung tidak ditemukan, contoh: {hilang['f_gabung'].head(2).tolist()}")

    arg.out.mkdir(parents=True, exist_ok=True)
    tmp = arg.out / "_gambar_sementara"
    tmp.mkdir(exist_ok=True)

    wb = Workbook()
    p = wb.active
    p.title = "Petunjuk"
    p.column_dimensions["A"].width = 26
    p.column_dimensions["B"].width = 110
    baris = [
        ("FORMULIR EVALUATOR SESI A: DESKRIPSI KONDISI BANGUNAN", None),
        ("Penelitian", "Ekstraksi Informasi Kerusakan Bangunan Pascabencana Menggunakan Small "
                       "Vision-Language Model pada Citra Satelit Resolusi Sangat Tinggi "
                       "(studi kasus gempa bumi dan tsunami Palu 2018)"),
        ("Peneliti", "Rizky Alif Ichwanto (222212857, 4SD2), Politeknik Statistika STIS"),
        ("Nama evaluator", ""),
        ("Kode evaluator", ""),
        ("Tanggal pengisian", ""),
        ("", ""),
        ("Tugas", "Anda melihat 40 bangunan. Setiap baris berisi satu citra dengan dua panel: "
                  "kiri SEBELUM bencana (24 Januari 2017), kanan SESUDAH bencana (1 Oktober 2018), "
                  "pada lokasi yang sama, skala 0,5 meter per piksel."),
        ("Bangunan sasaran", "Bangunan yang dinilai adalah bangunan yang ditunjuk BIDIK SILANG di "
                             "tengah kotak magenta. Bangunan lain hanya konteks. Ukuran kotak selalu "
                             "sama dan tidak menyatakan ukuran bangunan."),
        ("Yang ditulis", INSTRUKSI),
        ("Tidak terlihat", "Bila bangunan sasaran benar-benar tidak terlihat (tertutup awan, bayangan "
                           "pekat, atau pohon), tulis: bangunan tidak terlihat."),
        ("Penting", "Kerjakan sendiri tanpa berdiskusi dengan evaluator lain dan tanpa bantuan AI. "
                    "Tidak ada jawaban benar atau salah; yang dibutuhkan adalah pengamatan Anda "
                    "sendiri. Perkiraan waktu 25-35 menit."),
        ("", ""),
        ("DAFTAR PEMANDU", None),
    ]
    for f in fitur:
        baris.append((f["label"], f["definisi"]))
    for i, (x, y) in enumerate(baris, 1):
        c = p.cell(i, 1, x)
        c.alignment = ATAS
        c.font = Font(bold=True, size=13 if i == 1 else 11) if (y is None or x in (
            "Nama evaluator", "Kode evaluator", "Tanggal pengisian")) else Font()
        if y is not None:
            c = p.cell(i, 2, y)
            c.alignment = ATAS
            if x in ("Nama evaluator", "Kode evaluator", "Tanggal pengisian"):
                c.fill, c.border = ISI, KOTAK
            p.row_dimensions[i].height = max(18, 15 * (1 + len(str(y)) // 105))

    w = wb.create_sheet("Deskripsi")
    kepala = ["No", "Kode", "Citra (kiri: sebelum | kanan: sesudah)",
              "Deskripsi kondisi bangunan sasaran (satu kalimat)", "Catatan (opsional)"]
    for j, (h, l) in enumerate(zip(kepala, [5, 7, 76, 70, 24]), 1):
        c = w.cell(1, j, h)
        c.font, c.fill, c.border = Font(bold=True), KEPALA, KOTAK
        c.alignment = Alignment(wrap_text=True, vertical="center")
        w.column_dimensions[c.column_letter].width = l
    w.row_dimensions[1].height = 30
    w.freeze_panes = "D2"

    for i, r in enumerate(s.itertuples(index=False), 2):
        im = Image.open(r.f_gabung).convert("RGB")
        tinggi = round(im.height * LEBAR_GAMBAR / im.width)
        f_kecil = tmp / f"{r.kode}.png"
        im.resize((LEBAR_GAMBAR, tinggi), Image.LANCZOS).save(f_kecil)
        w.add_image(XLImage(str(f_kecil)), f"C{i}")
        w.row_dimensions[i].height = tinggi * 0.75 + 6
        w.cell(i, 1, i - 1)
        w.cell(i, 2, r.kode).font = Font(bold=True)
        for j in range(1, len(kepala) + 1):
            w.cell(i, j).border = KOTAK
            w.cell(i, j).alignment = Alignment(wrap_text=True, vertical="center")
        w.cell(i, 4).fill = ISI
        w.cell(i, 4).alignment = ATAS

    for ws in (p, w):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    w.print_title_rows = "1:1"
    wb.active = 0

    f_tpl = arg.out / "Formulir Sesi A - TEMPLATE.xlsx"
    wb.save(f_tpl)
    for n in range(1, arg.jumlah + 1):
        shutil.copy(f_tpl, arg.out / f"Formulir Sesi A - EVAL{n:02d}.xlsx")
    for f in tmp.glob("*.png"):
        f.unlink()
    tmp.rmdir()
    s[["kode", "uid", "tile", "kelas_xbd", "chip_pra", "chip_pasca"]].to_csv(
        arg.out / "kunci_formulir_sesi_a.csv", index=False, encoding="utf-8-sig")

    print("=" * 72)
    print(" FORMULIR EVALUATOR SESI A")
    print("=" * 72)
    print(f" Folder     : {arg.out}")
    print(f" Berkas     : Formulir Sesi A - EVAL01.xlsx ... EVAL{arg.jumlah:02d}.xlsx")
    print(f" Bangunan   : {len(s)} (urutan diacak, benih {BENIH}, kode E01-E{len(s):02d})")
    print(f" Pemandu    : {', '.join(f['label'] for f in fitur)}")
    print(" kunci_formulir_sesi_a.csv memuat label xBD: JANGAN dibagikan.")
    print(" Catat evaluator mana memegang kode EVAL berapa.")


if __name__ == "__main__":
    main()
