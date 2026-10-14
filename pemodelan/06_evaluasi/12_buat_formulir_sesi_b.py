#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
12_buat_formulir_sesi_b.py
==========================
Formulir evaluator mahasiswa SESI B (tujuan 2, evaluasi manusia).

Evaluator menilai keluaran tiga model pada 40 bangunan sampel kecil dengan tiga
kriteria QUEST (Tam dkk., 2024): accuracy, relevance, satisfaction; skala Likert
1-5. Keandalan antarevaluator dihitung dengan Krippendorff's alpha di skrip analisis.

  * Kode bangunan E01-E40 dan urutannya SAMA dengan sesi A (dibaca dari
    kunci_formulir_sesi_a.csv).
  * Nama model DISAMARKAN menjadi Model A/B/C, urutannya diacak per bangunan
    (benih 42), agar penilaian tidak dipengaruhi merek model. Pemetaan huruf ->
    model disimpan di kunci_formulir_sesi_b.csv, hanya untuk peneliti.
  * Yang ditampilkan per model: kelas kerusakan, 4 fitur, dan deskripsi, diambil
    dari hasil urai terbaru (hasil_terurai_p<sha>.csv bila ada).

Syarat: dibagikan SETELAH formulir sesi A kembali.

Penggunaan (dari folder KOMPONEN_4):
    python 12_buat_formulir_sesi_b.py                      # pilih run otomatis
    python 12_buat_formulir_sesi_b.py --runs <run1> <run2> <run3>
Pemilihan otomatis: untuk setiap model, run berstatus 'selesai', bukan uji cepat,
yang mencakup ke-40 bangunan sampel kecil, dan paling baru.
"""

from __future__ import annotations

import argparse
import json
import os
import random
import shutil
import sys
from pathlib import Path

try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from PIL import Image
except ImportError:  # pragma: no cover
    sys.exit("Butuh pandas, openpyxl, pillow. Jalankan:  pip install pandas openpyxl pillow")

AKAR = Path(os.environ.get("SKRIPSI_AKAR",
             r"F:\KULIAH STIS\TINGKAT IV - SD\SKRIPSWEET\MANUVER SKRIPSI AFTER SENSUS"))
RUNS = AKAR / "KOMPONEN_5" / "runs"
GABUNG = AKAR / "KOMPONEN_3" / "Output Chip" / "gabung"
FORMULIR = AKAR / "KOMPONEN_6" / "Formulir"
MODEL_DIHARAPKAN = ["qwen2.5-vl-3b", "internvl3-2b"]   # + satu model komersial
BENIH = 42
LEBAR_GAMBAR = 440
HURUF = "ABCDEFG"

KRITERIA = [
    ("Akurasi", "Informasi pada keluaran ini sesuai dengan kondisi bangunan sasaran yang tampak "
                "pada citra."),
    ("Relevansi", "Keluaran ini membahas bangunan sasaran dan aspek kerusakan yang ditanyakan, "
                  "tanpa informasi yang menyimpang atau tidak berhubungan."),
    ("Kepuasan", "Secara keseluruhan, saya puas dengan keluaran ini sebagai informasi kerusakan "
                 "bangunan untuk keperluan penanganan pascabencana."),
]
SKALA = "1 = sangat tidak setuju, 2 = tidak setuju, 3 = netral, 4 = setuju, 5 = sangat setuju"
TIPIS = Side(style="thin", color="999999")
TEBAL_SISI = Side(style="medium", color="444444")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
ISI = PatternFill("solid", fgColor="FFF7D6")
KEPALA = PatternFill("solid", fgColor="DCE6F1")


def hasil_run(d: Path) -> pd.DataFrame:
    calon = sorted(d.glob("hasil_terurai_p*.csv"), key=lambda f: f.stat().st_mtime)
    f = calon[-1] if calon else d / "hasil_terurai.csv"
    if not f.exists():
        sys.exit(f"{d.name}: belum ada hasil_terurai. Jalankan 10_rekap_run.py dulu.")
    return pd.read_csv(f).drop_duplicates("uid").set_index("uid")


def pilih_otomatis(uid40: set) -> list[str]:
    calon = []
    for d in RUNS.iterdir():
        f = d / "manifest.json"
        if not f.exists():
            continue
        m = json.loads(f.read_text(encoding="utf-8"))
        if m.get("jenis") in ("pengembangan", "kewajaran") or m.get("status") != "selesai" or m.get("uji_cepat"):
            continue
        # Sesi B menilai tiga model zero-shot. Model hasil fine-tuning (LoRA) hanya dilatih
        # untuk kelas, sehingga dievaluasi lewat confusion matrix dan evaluasi otomatis saja.
        if "lora" in m["model"]["nama"].lower():
            continue
        try:
            h = hasil_run(d)
        except SystemExit:
            continue
        if uid40 <= set(h.index):
            calon.append((m["model"]["nama"], m.get("waktu_mulai_utc", ""), d.name))
    terbaru = {}
    for model, waktu, nama in sorted(calon, key=lambda x: x[1]):
        terbaru[model] = nama
    return list(terbaru.values())


def teks_keluaran(r, fitur) -> str:
    if pd.isna(r.get("kelas_kerusakan")) and not r.get("f_sah", False):
        return "(keluaran model tidak dapat dibaca)"
    baris = [f"Kelas kerusakan: {str(r.get('kelas_kerusakan') or '-').replace('_', ' ')}"]
    for f in fitur:
        v = r.get(f"fitur_{f['kunci']}")
        baris.append(f"{f['label']}: {str(v).replace('_', ' ') if pd.notna(v) else '-'}")
    d = r.get("deskripsi")
    baris.append(f"Deskripsi: {d if isinstance(d, str) and d.strip() else '-'}")
    return "\n".join(baris)


def main() -> None:
    a = argparse.ArgumentParser()
    a.add_argument("--runs", nargs="+", help="nama folder run, satu per model (opsional)")
    a.add_argument("--jumlah", type=int, default=5)
    arg = a.parse_args()

    f_kunci_a = FORMULIR / "kunci_formulir_sesi_a.csv"
    if not f_kunci_a.exists():
        sys.exit("kunci_formulir_sesi_a.csv belum ada. Jalankan 11_buat_formulir_sesi_a.py dulu.")
    s = pd.read_csv(f_kunci_a)
    runs = arg.runs or pilih_otomatis(set(s["uid"]))

    model = []
    for nama in runs:
        d = RUNS / nama
        m = json.loads((d / "manifest.json").read_text(encoding="utf-8"))
        h = hasil_run(d)
        kurang = set(s["uid"]) - set(h.index)
        if kurang:
            sys.exit(f"{nama}: {len(kurang)} dari 40 bangunan tidak ada di run ini.")
        model.append({"run": nama, "model": m["model"]["nama"], "hasil": h,
                      "konfig": json.loads((d / "paket" / "konfigurasi_fitur.json").read_text(encoding="utf-8"))})
    nama_model = [m["model"] for m in model]
    if len(set(nama_model)) != len(nama_model):
        sys.exit(f"Ada dua run dari model yang sama: {nama_model}. Pakai --runs.")
    if not all(x in nama_model for x in MODEL_DIHARAPKAN) or len(model) != 3:
        sys.exit(f"Diharapkan tepat 3 model (2 SVLM + 1 komersial), ditemukan: {nama_model}. Pakai --runs.")

    k = model[0]["konfig"]
    peta = {c["kunci"]: c for c in k["kandidat"]}
    fitur = [peta[t] for t in k["seleksi"]["terpilih"]]
    rng = random.Random(BENIH)
    tmp = FORMULIR / "_gambar_sementara_b"
    tmp.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    p = wb.active
    p.title = "Petunjuk"
    p.column_dimensions["A"].width = 22
    p.column_dimensions["B"].width = 110
    petunjuk = [
        ("FORMULIR EVALUATOR SESI B: PENILAIAN KELUARAN MODEL", None),
        ("Nama evaluator", ""), ("Kode evaluator", "(sama dengan sesi A, misal EVAL01)"),
        ("Tanggal pengisian", ""), ("", ""),
        ("Tugas", "Untuk setiap bangunan (E01-E40, sama seperti sesi A) terdapat keluaran dari tiga "
                  "model, disebut Model A, B, dan C. Nilailah SETIAP keluaran pada tiga kriteria di "
                  "bawah dengan angka 1-5. Nama model sengaja disamarkan dan urutannya diacak."),
        ("Bangunan sasaran", "Bangunan yang ditunjuk bidik silang di tengah kotak magenta. "
                             "Citra kiri sebelum bencana, kanan sesudah bencana."),
        ("Skala", SKALA),
    ] + [(n, x) for n, x in KRITERIA] + [
        ("", ""),
        ("Catatan", "Nilailah setiap keluaran secara mandiri, walaupun beberapa keluaran tampak sama. "
                    "Kerjakan sendiri, tanpa berdiskusi dan tanpa bantuan AI. Perkiraan waktu 45-60 menit."),
    ]
    for i, (x, y) in enumerate(petunjuk, 1):
        c = p.cell(i, 1, x)
        c.font = Font(bold=True, size=13 if i == 1 else 11)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if y is not None:
            c = p.cell(i, 2, y)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if x in ("Nama evaluator", "Kode evaluator", "Tanggal pengisian"):
                c.fill, c.border = ISI, KOTAK
            p.row_dimensions[i].height = max(18, 15 * (1 + len(str(y)) // 105))

    w = wb.create_sheet("Penilaian")
    kepala = (["No", "Kode", "Citra (kiri: sebelum | kanan: sesudah)", "Model", "Keluaran model"]
              + [f"{n} (1-5)" for n, _ in KRITERIA] + ["Catatan (opsional)"])
    for j, (h, l) in enumerate(zip(kepala, [5, 7, 64, 8, 62, 12, 12, 12, 22]), 1):
        c = w.cell(1, j, h)
        c.font, c.fill, c.border = Font(bold=True), KEPALA, KOTAK
        c.alignment = Alignment(wrap_text=True, vertical="center")
        w.column_dimensions[c.column_letter].width = l
    w.row_dimensions[1].height = 32
    w.freeze_panes = "F2"
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Nilai tidak sah", error="Isi angka bulat 1 sampai 5.")
    w.add_data_validation(dv)

    kunci_b, baris = [], 2
    for no, r in enumerate(s.itertuples(index=False), 1):
        urut = list(range(len(model)))
        rng.shuffle(urut)
        im = Image.open(GABUNG / str(r.chip_pra).replace("_pra", "_gabung")).convert("RGB")
        tinggi = round(im.height * LEBAR_GAMBAR / im.width)
        f_kecil = tmp / f"{r.kode}.png"
        im.resize((LEBAR_GAMBAR, tinggi), Image.LANCZOS).save(f_kecil)
        for slot, idx in enumerate(urut):
            mdl = model[idx]
            w.cell(baris, 4, f"Model {HURUF[slot]}").font = Font(bold=True)
            w.cell(baris, 5, teks_keluaran(mdl["hasil"].loc[r.uid], fitur))
            for j in range(1, len(kepala) + 1):
                w.cell(baris, j).border = KOTAK
                w.cell(baris, j).alignment = Alignment(wrap_text=True, vertical="top")
            for j in range(6, 6 + len(KRITERIA)):
                w.cell(baris, j).fill = ISI
                dv.add(w.cell(baris, j))
            w.row_dimensions[baris].height = max(120, (tinggi * 0.75 + 8) / len(model))
            kunci_b.append({"kode": r.kode, "uid": r.uid, "kelas_xbd": r.kelas_xbd,
                            "huruf": HURUF[slot], "model": mdl["model"], "run": mdl["run"],
                            "baris_excel": baris})
            baris += 1
        awal = baris - len(model)
        for kol in (1, 2, 3):
            w.merge_cells(start_row=awal, start_column=kol, end_row=baris - 1, end_column=kol)
        w.cell(awal, 1, no).alignment = Alignment(vertical="top")
        w.cell(awal, 2, r.kode).font = Font(bold=True)
        w.cell(awal, 2).alignment = Alignment(vertical="top")
        w.add_image(XLImage(str(f_kecil)), f"C{awal}")
        for j in range(1, len(kepala) + 1):
            w.cell(baris - 1, j).border = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TEBAL_SISI)

    for ws in (p, w):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    w.print_title_rows = "1:1"

    f_tpl = FORMULIR / "Formulir Sesi B - TEMPLATE.xlsx"
    wb.save(f_tpl)
    for n in range(1, arg.jumlah + 1):
        shutil.copy(f_tpl, FORMULIR / f"Formulir Sesi B - EVAL{n:02d}.xlsx")
    for f in tmp.glob("*.png"):
        f.unlink()
    tmp.rmdir()
    pd.DataFrame(kunci_b).to_csv(FORMULIR / "kunci_formulir_sesi_b.csv", index=False,
                                 encoding="utf-8-sig")
    print("=" * 72)
    print(" FORMULIR EVALUATOR SESI B (QUEST)")
    print("=" * 72)
    for m in model:
        print(f"  {m['model']:<18} <- {m['run']}")
    print(f" {len(s)} bangunan x {len(model)} model = {len(kunci_b)} keluaran per evaluator")
    print(f" Berkas: {FORMULIR}  (EVAL01-EVAL{arg.jumlah:02d})")
    print(" kunci_formulir_sesi_b.csv memetakan Model A/B/C ke model asli: JANGAN dibagikan.")


if __name__ == "__main__":
    main()
