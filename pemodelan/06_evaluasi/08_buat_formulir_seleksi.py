#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
08_buat_formulir_seleksi.py
===========================
Membuat formulir penilaian kepentingan kandidat fitur dari konfigurasi_fitur.json.

Formulir dibangkitkan, tidak diketik manual, supaya definisi dan tingkatan nilai
di formulir selalu sama dengan yang dipakai pada prompt model dan formulir evaluator.

Pengisi hanya mengisi dua kolom: skor Likert 1-5 dan alasan singkat.
Peringkat dihitung otomatis dengan rumus RANK, jadi tidak perlu diisi manual.

    python 08_buat_formulir_seleksi.py
    python 08_buat_formulir_seleksi.py --peran peneliti
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("Butuh openpyxl. Jalankan:  pip install openpyxl")

KONFIG = Path(__file__).resolve().parent / "konfigurasi_fitur.json"
OUT = Path(__file__).resolve().parent.parent / "KOMPONEN_1"

KUNING = PatternFill("solid", fgColor="FFF2CC")
BIRU = PatternFill("solid", fgColor="DDE3EA")
ABU = PatternFill("solid", fgColor="F5F5F5")
HIJAU = PatternFill("solid", fgColor="E2EFDA")
TEBAL = Font(bold=True)
KECIL = Font(size=9)
GARIS = Border(*[Side(style="thin", color="BFBFBF")] * 4)
ATAS = Alignment(vertical="top", wrap_text=True)


def petunjuk(ws, k, peran, nama, tanggal):
    n = k["seleksi"]["jumlah_terpilih"]
    ws.column_dimensions["A"].width = 118
    baris = [
        ("FORMULIR PENILAIAN KEPENTINGAN KANDIDAT FITUR", "judul"),
        ("Ekstraksi Informasi Kerusakan Bangunan Pascabencana Menggunakan Small Vision-Language Model", None),
        ("pada Citra Satelit Resolusi Sangat Tinggi. Studi kasus: Gempa Bumi dan Tsunami Palu 2018", None),
        ("", None),
        ("Peneliti       : Rizky Alif Ichwanto (222212857, 4SD2)", None),
        ("Program studi  : D-IV Komputasi Statistik, Politeknik Statistika STIS", None),
        (f"Peran pengisi  : {peran}", None),
        (f"Nama pengisi   : {nama}", None),
        (f"Tanggal        : {tanggal}", None),
        ("", None),
        ("A.  Tujuan", "bab"),
        ("Menyeleksi kandidat fitur visual kerusakan bangunan yang akan dipakai sebagai dasar penentuan", None),
        ("kelas kerusakan. Fitur terpilih dituangkan ke dalam prompt model dan ke dalam formulir penilaian", None),
        ("evaluator. Penilaian ini menjawab tujuan penelitian ke-1.", None),
        ("", None),
        ("B.  Data yang dinilai", "bab"),
        ("Sumber citra    : dataset xBD/xView2, sensor WorldView-02, pasangan citra pra dan pasca bencana", None),
        ("Pasangan citra  : pra 24 Januari 2017, pasca 1 Oktober 2018. Dikunci pada satu pasangan akuisisi", None),
        ("Skala           : 0,5 meter per piksel", None),
        ("Unit analisis   : satu bangunan, disajikan sebagai potongan citra 128 x 128 piksel (64 x 64 meter),", None),
        ("                  diperbesar ke 448 x 448 piksel", None),
        ("Penanda         : bangunan yang dinilai ditandai kotak magenta dan bidik silang di titik pusatnya", None),
        ("Ukuran bangunan : median sisi terpanjang 29 piksel (sekitar 14 meter); 76 persen di bawah 40 piksel", None),
        ("", None),
        ("C.  Cara mengisi", "bab"),
        ('Buka sheet "Penilaian Fitur". Isi dua kolom berlatar kuning saja.', None),
        ("", None),
        ("   1.  Kolom H, Skor Kepentingan. Angka 1 sampai 5.", None),
        ("          1 = sangat tidak penting      2 = tidak penting      3 = cukup penting", None),
        ("          4 = penting                   5 = sangat penting", None),
        ("", None),
        ("   2.  Kolom I, Alasan singkat. Satu sampai dua kalimat. Alasan ini dikutip sebagai justifikasi", None),
        ("       pemilihan fitur di dalam buku skripsi.", None),
        ("", None),
        ("Peringkat terisi otomatis dari skor, jadi tidak perlu diisi. Sheet Rekapitulasi juga otomatis.", None),
        ("", None),
        (f"D.  Jumlah fitur yang dipakai: {n}", "bab"),
        (f"Hanya {n} fitur dengan skor tertinggi yang dipakai. Angka {n} adalah ketetapan peneliti, bukan", None),
        ("bagian dari penilaian ini. Dasarnya dua hal:", None),
        ("", None),
        ("   1.  Batas teknis model. Pengujian 21 September menunjukkan model berukuran 2 sampai 3 miliar", None),
        ("       parameter berhenti memberi jawaban yang membedakan ketika diminta menilai enam fitur", None),
        ("       sekaligus dalam satu prompt.", None),
        ("   2.  Beban evaluator. Setiap evaluator menilai 40 potongan citra. Makin banyak fitur, makin lama", None),
        ("       waktu pengisian dan makin turun konsistensi penilaiannya.", None),
        ("", None),
        ("E.  Bila ada nilai seri pada batas", "bab"),
        ("Peringkat dihitung dari skor, sehingga skor yang sama menghasilkan peringkat yang sama. Bila jumlah", None),
        (f"fitur yang masuk {n} besar melebihi {n} karena seri, sheet Rekapitulasi akan menandainya, dan peneliti", None),
        ("akan menghubungi pengisi untuk menentukan urutan di antara yang seri saja.", None),
        ("", None),
        ("F.  Keterangan kolom pada sheet Penilaian Fitur", "bab"),
        ("Kolom E  : nilai yang mungkin untuk tiap fitur, ditulis berurutan dari kondisi paling ringan ke", None),
        ('           paling berat. Nilai "tidak_jelas" dipakai bila bangunan tidak teramati sama sekali,', None),
        ("           dan bukan merupakan tingkat kerusakan.", None),
        ("Kolom F  : sumber rujukan tingkatan nilai tersebut.", None),
        ("Kolom G  : catatan hasil pemeriksaan metadata citra. Catatan ini informasi, bukan arahan penilaian.", None),
    ]
    for i, (teks, jenis) in enumerate(baris, 1):
        sel = ws.cell(row=i, column=1, value=teks)
        if jenis == "judul":
            sel.font = Font(bold=True, size=14, color="1F3864")
        elif jenis == "bab":
            sel.font = Font(bold=True, size=11, color="2E5496")
        else:
            sel.font = Font(size=10.5)
    ws.sheet_view.showGridLines = False


def penilaian(ws, k, peran):
    kand = k["kandidat"]
    n0, n1 = 5, 4 + len(kand)
    ws["A1"] = f"PENILAIAN KEPENTINGAN KANDIDAT FITUR — diisi oleh {peran}"
    ws["A1"].font = Font(bold=True, size=12, color="1F3864")
    ws["A2"] = ("Isi kolom H (skor 1-5) dan kolom I (alasan) saja. Kolom J terisi otomatis dari kolom H.")
    ws["A2"].font = Font(size=10)

    judul = ["No", "Kelompok", "Kandidat Fitur", "Definisi Operasional",
             "Nilai yang Mungkin (urut ringan ke berat)", "Sumber Rujukan Tingkatan",
             "Catatan Teknis dari Metadata Citra", "Skor Kepentingan (1-5)",
             "Alasan Singkat", "Peringkat (otomatis)"]
    lebar = [4, 16, 26, 44, 34, 46, 44, 11, 36, 11]
    for j, (t, w) in enumerate(zip(judul, lebar), 1):
        s = ws.cell(row=4, column=j, value=t)
        s.font = TEBAL; s.fill = BIRU; s.border = GARIS
        s.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 32

    for i, c in enumerate(kand):
        r = n0 + i
        nilai = " > ".join(c["level"]) + '   |   "tidak_jelas" bila tidak teramati'
        isi = [c["no"], c["kelompok"], c["label"], c["definisi"], nilai,
               c.get("rujukan_level", ""), c.get("keteramatan", ""), None, None,
               f'=IF($H{r}="","",RANK($H{r},$H${n0}:$H${n1},0))']
        for j, v in enumerate(isi, 1):
            s = ws.cell(row=r, column=j, value=v)
            s.border = GARIS; s.alignment = ATAS
            s.font = KECIL if j in (4, 6, 7) else Font(size=10)
            if j in (8, 9):
                s.fill = KUNING
            elif j == 10:
                s.fill = ABU
                s.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 76

    dv = DataValidation(type="whole", operator="between", formula1=1, formula2=5,
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Skor tidak sah", error="Isi angka 1 sampai 5 saja.")
    ws.add_data_validation(dv); dv.add(f"H{n0}:H{n1}")

    r = n1 + 3
    ws.cell(row=r, column=3, value="Nama dan tanda tangan pengisi:").font = Font(size=10)
    ws.cell(row=r + 4, column=3, value="(............................................)").font = Font(size=10)
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def rekap(ws, k):
    kand = k["kandidat"]
    n = k["seleksi"]["jumlah_terpilih"]
    n0, n1 = 5, 4 + len(kand)
    ws["A1"] = "REKAPITULASI HASIL SELEKSI (terisi otomatis, tidak perlu diisi)"
    ws["A1"].font = Font(bold=True, size=12, color="1F3864")
    ws["A2"] = (f"Peringkat dihitung dari skor Likert. Skor yang sama menghasilkan peringkat yang sama. "
                f"Fitur dengan peringkat {n} ke atas dinyatakan TERPILIH.")
    ws["A2"].font = Font(size=10)

    judul = ["No", "Kandidat Fitur", "Skala Ukur", "Skor", "Peringkat", "Status"]
    lebar = [4, 40, 16, 8, 10, 16]
    for j, (t, w) in enumerate(zip(judul, lebar), 1):
        s = ws.cell(row=4, column=j, value=t)
        s.font = TEBAL; s.fill = BIRU; s.border = GARIS
        s.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w

    for i, c in enumerate(kand):
        r = n0 + i
        isi = [c["no"], c["label"], c["skala_ukur"],
               f"='Penilaian Fitur'!H{r}",
               f'=IF($D{r}="","",RANK($D{r},$D${n0}:$D${n1},0))',
               f'=IF($D{r}="","",IF($E{r}<=$H$2,"TERPILIH","tidak dipakai"))']
        for j, v in enumerate(isi, 1):
            s = ws.cell(row=r, column=j, value=v)
            s.border = GARIS
            s.alignment = Alignment(horizontal="center" if j in (1, 4, 5, 6) else "left",
                                    vertical="center", wrap_text=j == 2)
            s.font = Font(size=10)

    ws["G2"] = "Jumlah fitur yang dipakai:"; ws["G2"].font = TEBAL
    ws["H2"] = n; ws["H2"].font = Font(bold=True, size=12); ws["H2"].fill = HIJAU
    ws["H2"].alignment = Alignment(horizontal="center")
    ws.column_dimensions["G"].width = 30; ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 62

    cek = [
        ("Jumlah kandidat yang sudah dinilai", f"=COUNT(D{n0}:D{n1})",
         f'=IF(COUNT(D{n0}:D{n1})={len(kand)},"lengkap","BELUM LENGKAP, masih ada yang kosong")'),
        ("Jumlah fitur berstatus TERPILIH", f'=COUNTIF(F{n0}:F{n1},"TERPILIH")',
         f'=IF(COUNTIF(F{n0}:F{n1},"TERPILIH")=0,"belum ada skor",'
         f'IF(COUNTIF(F{n0}:F{n1},"TERPILIH")>$H$2,'
         f'"ADA NILAI SERI DI BATAS. Peneliti akan menghubungi pengisi untuk menentukan urutannya.",'
         f'"sesuai, tidak ada seri di batas"))'),
        ("Rata-rata skor seluruh kandidat", f"=IFERROR(AVERAGE(D{n0}:D{n1}),\"\")", ""),
    ]
    for i, (label, rumus, ket) in enumerate(cek):
        r = n1 + 2 + i
        ws.cell(row=r, column=2, value=label).font = TEBAL
        ws.cell(row=r, column=4, value=rumus).alignment = Alignment(horizontal="center")
        if ket:
            s = ws.cell(row=r, column=6, value=ket)
            s.font = Font(size=10); s.alignment = Alignment(horizontal="left")
    ws.sheet_view.showGridLines = False


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--konfig", type=Path, default=KONFIG)
    p.add_argument("--out", type=Path, default=OUT)
    p.add_argument("--peran", choices=["dosen", "peneliti"], default="dosen")
    a = p.parse_args()

    k = json.loads(a.konfig.read_text(encoding="utf-8"))
    if a.peran == "dosen":
        peran = "ANNOTATOR (ahli penginderaan jauh / geospasial)"
        nama, tgl = "______________________________", "______________________________"
        berkas = "Formulir Seleksi Fitur Penting - ANNOTATOR.xlsx"
    else:
        peran = "PENILAI KEDUA (peneliti)"
        nama, tgl = "Rizky Alif Ichwanto (222212857, 4SD2)", "______________________________"
        berkas = "Formulir Seleksi Fitur Penting - PENILAI KEDUA.xlsx"

    wb = Workbook()
    petunjuk(wb.active, k, peran, nama, tgl); wb.active.title = "Petunjuk"
    penilaian(wb.create_sheet("Penilaian Fitur"), k, peran)
    rekap(wb.create_sheet("Rekapitulasi"), k)

    a.out.mkdir(parents=True, exist_ok=True)
    f = a.out / berkas
    wb.save(f)

    print("=" * 74)
    print(" FORMULIR SELEKSI FITUR SIAP")
    print("=" * 74)
    print(f" Peran pengisi   : {peran}")
    print(f" Berkas          : {f}")
    print(f" Kandidat        : {len(k['kandidat'])}")
    print(f" Fitur dipakai   : {k['seleksi']['jumlah_terpilih']} teratas")
    n_ord = sum(1 for c in k["kandidat"] if c["skala_ukur"] == "ordinal")
    print(f" Skala ukur      : {n_ord} ordinal, {len(k['kandidat']) - n_ord} biner")
    print(f" Yang diisi      : kolom H (skor 1-5) dan kolom I (alasan). Peringkat otomatis.")
    print("=" * 74)


if __name__ == "__main__":
    main()
