#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
17_buat_formulir_prototipe.py
=============================
Formulir uji prototipe (tujuan 3) untuk 5 evaluator mahasiswa (pengguna proksi):
  (1) Black-box testing: skenario disusun per fungsi dengan partisi ekuivalensi (Bab III 3.2.6);
      penguji menuliskan keluaran yang muncul dan status Berhasil/Gagal.
  (2) SUS versi bahasa Indonesia (Sharfina dan Santoso, 2016), 10 butir, skala 1-5, diisi SETELAH skenario.
Keluaran yang diharapkan dihitung dari data prototipe terbaru (KOMPONEN_8/prototipe_*/data/data.js),
sehingga angka pada formulir selalu sama dengan aplikasi yang diuji.

Keluaran: KOMPONEN_6/Formulir/Formulir Uji Prototipe - EVAL01..EVAL05.xlsx

Penggunaan (dari folder KOMPONEN_4, setelah 16_bangun_prototipe.py):
    python 17_buat_formulir_prototipe.py
"""

from __future__ import annotations

import json
import os
import shutil
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("Butuh openpyxl. Jalankan:  pip install openpyxl")

AKAR = Path(os.environ.get("SKRIPSI_AKAR",
            r"F:\KULIAH STIS\TINGKAT IV - SD\SKRIPSWEET\MANUVER SKRIPSI AFTER SENSUS"))
PROTO = AKAR / "KOMPONEN_8"
FORMULIR = AKAR / "KOMPONEN_6" / "Formulir"
JUMLAH = 5

SUS = [  # Sharfina dan Santoso (2016), versi akhir adaptasi bahasa Indonesia
    "Saya berpikir akan menggunakan sistem ini lagi.",
    "Saya merasa sistem ini rumit untuk digunakan.",
    "Saya merasa sistem ini mudah untuk digunakan.",
    "Saya membutuhkan bantuan dari orang lain atau teknisi dalam menggunakan sistem ini.",
    "Saya merasa fitur-fitur sistem ini berjalan dengan semestinya.",
    "Saya merasa ada banyak hal yang tidak konsisten (tidak serasi) pada sistem ini.",
    "Saya merasa orang lain akan memahami cara menggunakan sistem ini dengan cepat.",
    "Saya merasa sistem ini membingungkan.",
    "Saya merasa tidak ada hambatan dalam menggunakan sistem ini.",
    "Saya perlu membiasakan diri terlebih dahulu sebelum menggunakan sistem ini.",
]
TIPIS = Side(style="thin", color="999999")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
ISI = PatternFill("solid", fgColor="FFF7D6")
KEPALA = PatternFill("solid", fgColor="DCE6F1")
ATAS = Alignment(wrap_text=True, vertical="top")


def muat_data() -> tuple[dict, Path]:
    calon = sorted(p for p in PROTO.glob("prototipe_*") if p.is_dir())
    if not calon:
        sys.exit("Prototipe belum dibangun. Jalankan 16_bangun_prototipe.py dulu.")
    teks = (calon[-1] / "data" / "data.js").read_text(encoding="utf-8")
    return json.loads(teks[teks.index("{"): teks.rindex("}") + 1]), calon[-1]


def skenario(D: dict) -> list[tuple]:
    B = D["bangunan"]
    per = {b["kode"]: b for b in B}
    n = len(B)
    h = {k: sum(b["kelas"] == k for b in B) for k in ("tidak_rusak", "rusak", "hancur")}
    nm = {"tidak_rusak": "tidak rusak", "rusak": "rusak", "hancur": "hancur"}
    xbd = {"tidak_rusak": "no-damage", "rusak": "major-damage", "hancur": "destroyed"}
    tr_beda = sum(b["kelas"] == "tidak_rusak" and b["xbd"] != "tidak_rusak" for b in B)
    frek = {}
    for b in B:
        frek[b["deskripsi"]] = frek.get(b["deskripsi"], 0) + 1
    k1, k2 = "PL-015", "PL-200"
    b1 = per[k1]
    seragam = (f" Di bawah deskripsi muncul catatan bahwa kalimat yang sama dipakai untuk {frek[b1['deskripsi']]} bangunan."
               if frek.get(b1["deskripsi"], 0) > 1 else "")
    atap1 = D["label_level"]["keutuhan_atap"].get(b1["fitur"]["keutuhan_atap"] or "", "tidak terbaca")
    return [
        # (kode, fungsi, kelas partisi, langkah, masukan, keluaran yang diharapkan)
        ("B01", "Informasi Umum", "Membaca informasi umum",
         "Buka prototipe. Pada tab Informasi Umum, gulir dari atas sampai bawah.", "-",
         "Enam bagian tampil berurutan (Kejadian, Data, Cara kerja, Hasil, Batasan, Tentang prototipe) beserta ikon "
         "pada angka korban, tiga contoh citra, dan satu tabel kesesuaian."),
        ("B02", "Navigasi dan ringkasan", "Pindah halaman", "Di akhir halaman, klik tombol \"Buka Dashboard\". "
         "Amati kotak ringkasan sebelum memakai penyaring.", "klik tombol",
         f"Tab Dashboard terbuka. Bangunan ditampilkan {n}; kelas model tidak rusak {h['tidak_rusak']}, rusak {h['rusak']}, "
         f"hancur {h['hancur']}."),
        ("B03", "Penyaring", "Satu penyaring (masukan sah)", "Pilih Kelas model = Hancur. Tidak perlu menekan tombol apa pun.",
         "Hancur", f"Ringkasan, peta, daftar, dan pilihan Kode bangunan langsung berubah dan hanya berisi bangunan kelas "
         f"hancur ({h['hancur']} bangunan)."),
        ("B04", "Penyaring", "Gabungan penyaring", "Tambahkan Kesesuaian = Tidak sesuai, lalu ubah Kelas model = Tidak rusak.",
         "Tidak rusak + Tidak sesuai", f"Tampil {tr_beda} bangunan; di daftar, kolom Label xBD tidak memuat no-damage."),
        ("B05", "Penyaring", "Mengembalikan keadaan awal", "Klik Atur ulang.", "klik tombol",
         f"Semua penyaring kembali ke Semua; bangunan ditampilkan {n}."),
        ("B06", "Peta", "Mengganti peta dasar", "Amati peta, lalu arahkan kursor ke ikon lapisan di pojok kanan atas peta "
         "dan pilih \"Citra dasar Esri\".", "pilih lapisan",
         "Peta jalan tampil di bawah titik bangunan; setelah diganti, citra dasar tampil dan titik tetap berada di tempatnya."),
        ("B07", "Peta", "Ringkasan per kotak grid", "Centang \"Tampilkan ringkasan per kotak grid (±1 km)\", lalu arahkan "
         "kursor ke salah satu kotak.", "centang",
         "Kotak berwarna muncul di atas peta; keterangan menampilkan jumlah bangunan per kelas pada kotak tersebut."),
        ("B08", "Rincian bangunan", "Memilih dari peta", "Klik salah satu titik pada peta.", "klik titik",
         "Panel rincian menampilkan kode, citra sebelum-sesudah, kelas model, label xBD, empat tanda kerusakan, dan deskripsi."),
        ("B09", "Rincian bangunan", "Memilih dari daftar kode", f"Pada penyaring Kode bangunan, pilih {k1}.", k1,
         f"Rincian {k1} tampil: kelas model {nm[b1['kelas']]}, label xBD {xbd[b1['xbd']]}.{seragam}"),
        ("B10", "Unduh data", "Mengunduh hasil penyaringan", "Pilih Kelas model = Rusak, lalu klik Unduh CSV. Buka berkasnya.",
         "Rusak + klik tombol", f"Berkas CSV terunduh dan berisi {h['rusak']} baris bangunan, semuanya berkelas model Rusak."),
        ("B11", "Tanya-jawab", "Pertanyaan bangunan tanpa memilih bangunan",
         "Muat ulang halaman, buka tab Tanya-Jawab. Sebelum memilih bangunan, klik \"Bagaimana kondisi atapnya?\" pada "
         "Rekomendasi Pertanyaan di ruang tanya-jawab.", "klik rekomendasi",
         "Sistem meminta memilih bangunan terlebih dahulu."),
        ("B12", "Tanya-jawab", "Pertanyaan baku", f"Pilih {k1} pada daftar Kode bangunan, lalu klik rekomendasi "
         "\"Apa kelas kerusakan bangunan ini?\".", k1,
         f"Sistem mengumumkan {k1} terpilih beserta citranya; jawaban kelas model {nm[b1['kelas']]} dan label xBD "
         f"{xbd[b1['xbd']]}, sama dengan rincian di Dashboard."),
        ("B13", "Tanya-jawab", "Pertanyaan baku dengan kata sendiri", "Ketik \"atapnya gimana?\" lalu Kirim.",
         "atapnya gimana?", f"Dikenali sebagai pertanyaan atap; jawaban keutuhan atap {k1}: {atap1}."),
        ("B14", "Tanya-jawab", "Kode di dalam pertanyaan", f"Ketik \"tampilkan citra {k2}\" lalu Kirim.", f"tampilkan citra {k2}",
         f"Bangunan terpilih berganti ke {k2} dan citranya tampil."),
        ("B15", "Tanya-jawab", "Kode yang tidak ada", "Ketik \"bagaimana atap PL-999?\" lalu Kirim.", "bagaimana atap PL-999?",
         f"Muncul pesan kode tidak ditemukan (kode tersedia PL-001 sampai PL-{n:03d}); aplikasi tetap dapat dipakai."),
        ("B16", "Tanya-jawab", "Pertanyaan di luar cakupan", "Ketik \"Siapa pemilik bangunan ini?\" lalu Kirim.",
         "Siapa pemilik bangunan ini?", "Sistem menyatakan informasi tersebut tidak tersedia dan menampilkan rekomendasi pertanyaan."),
        ("B17", "Tanya-jawab", "Pertanyaan umum", "Klik \"Lihat semua pertanyaan\", lalu klik \"Berapa banyak bangunan per "
         "kelas kerusakan?\".", "klik rekomendasi",
         f"Jumlah per kelas model sama dengan ringkasan Dashboard ({h['tidak_rusak']} / {h['rusak']} / {h['hancur']})."),
        ("B18", "Tanya-jawab", "Memilih bangunan dari pasangan citra", "Pada bagian \"Atau pilih dari pasangan citra\", klik "
         "Berikutnya, lalu klik salah satu pasangan citra.", "klik citra",
         "Pasangan citra diberi bingkai; daftar Kode bangunan dan keterangan bangunan terpilih berubah; sistem mengumumkan "
         "bangunan terpilih."),
        ("B19", "Navigasi antarfitur", "Dari tanya-jawab ke Dashboard", "Pada pesan sistem terakhir yang mengumumkan bangunan "
         "terpilih, klik \"Lihat bangunan ini di Dashboard\".", "klik tautan",
         "Tab Dashboard terbuka dan panel rincian menampilkan bangunan yang sama dengan pilihan di tanya-jawab."),
    ]


def main():
    D, folder = muat_data()
    sk = skenario(D)
    wb = Workbook()
    p = wb.active
    p.title = "Petunjuk"
    p.column_dimensions["A"].width = 24
    p.column_dimensions["B"].width = 105
    baris = [
        ("FORMULIR UJI PROTOTIPE (TUJUAN 3)", None),
        ("Penelitian", "Ekstraksi Informasi Kerusakan Bangunan Pascabencana Menggunakan Small Vision-Language Model pada "
                       "Citra Satelit Resolusi Sangat Tinggi"),
        ("Peneliti", "Rizky Alif Ichwanto (222212857, 4SD2), Politeknik Statistika STIS"),
        ("Nama evaluator", ""), ("Kode evaluator", "(sama dengan Sesi A dan B, misal EVAL01)"), ("Tanggal pengisian", ""),
        ("Peramban yang dipakai", "(misal Chrome / Edge)"), ("", ""),
        ("Persiapan", "Buka tautan prototipe yang dikirim peneliti dengan Chrome atau Edge di laptop yang tersambung internet. "
                      "Bila tautan tidak dapat dibuka, ekstrak zip prototipe lalu buka index.html. Buku pedoman penggunaan "
                      "boleh dibaca, tetapi coba jalankan skenario tanpa bantuan lebih dulu."),
        ("Bagian 1", "Sheet \"Black-box\": jalankan setiap skenario B01-B19 berurutan. Tuliskan keluaran yang muncul "
                     "dengan singkat, lalu pilih status Berhasil bila sesuai dengan keluaran yang diharapkan atau Gagal bila tidak."),
        ("Bagian 2", "Sheet \"SUS\": setelah semua skenario selesai, beri nilai 1-5 pada kesepuluh pernyataan berdasarkan "
                     "pengalaman Anda memakai prototipe. 1 = sangat tidak setuju, 5 = sangat setuju. Sebagian pernyataan "
                     "bernada positif dan sebagian negatif; baca dengan teliti. Semua butir wajib diisi."),
        ("Penting", "Kerjakan sendiri, tanpa berdiskusi dan tanpa bantuan AI. Kritik dan temuan kesalahan sangat membantu. "
                    "Perkiraan waktu 30-40 menit."),
        ("Versi prototipe", folder.name),
    ]
    for i, (x, y) in enumerate(baris, 1):
        c = p.cell(i, 1, x)
        c.font = Font(bold=True, size=13 if i == 1 else 11)
        c.alignment = ATAS
        if y is not None:
            c = p.cell(i, 2, y)
            c.alignment = ATAS
            if x in ("Nama evaluator", "Kode evaluator", "Tanggal pengisian", "Peramban yang dipakai"):
                c.fill, c.border = ISI, KOTAK
            p.row_dimensions[i].height = max(18, 15 * (1 + len(str(y)) // 100))

    w = wb.create_sheet("Black-box")
    kep = ["Kode", "Fungsi", "Kelas masukan (partisi)", "Langkah", "Masukan", "Keluaran yang diharapkan",
           "Keluaran yang muncul (isi)", "Status", "Catatan (opsional)"]
    for j, (h, l) in enumerate(zip(kep, [7, 15, 22, 40, 18, 44, 38, 12, 22]), 1):
        c = w.cell(1, j, h)
        c.font, c.fill, c.border, c.alignment = Font(bold=True), KEPALA, KOTAK, Alignment(wrap_text=True, vertical="center")
        w.column_dimensions[c.column_letter].width = l
    dv = DataValidation(type="list", formula1='"Berhasil,Gagal"', allow_blank=True)
    w.add_data_validation(dv)
    for i, s in enumerate(sk, 2):
        for j, v in enumerate(s, 1):
            c = w.cell(i, j, v)
            c.alignment, c.border = ATAS, KOTAK
        for j in (7, 8, 9):
            c = w.cell(i, j)
            c.border, c.alignment = KOTAK, ATAS
            if j in (7, 8):
                c.fill = ISI
        dv.add(w.cell(i, 8))
        w.row_dimensions[i].height = 62
    w.freeze_panes = "C2"

    s = wb.create_sheet("SUS")
    for j, (h, l) in enumerate(zip(["No", "Pernyataan", "Nilai (1-5)"], [6, 90, 14]), 1):
        c = s.cell(1, j, h)
        c.font, c.fill, c.border = Font(bold=True), KEPALA, KOTAK
        s.column_dimensions[c.column_letter].width = l
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                         showErrorMessage=True, errorTitle="Nilai tidak sah", error="Isi angka bulat 1 sampai 5.")
    s.add_data_validation(dv2)
    for i, t in enumerate(SUS, 2):
        s.cell(i, 1, i - 1).border = KOTAK
        c = s.cell(i, 2, t)
        c.border, c.alignment = KOTAK, ATAS
        c = s.cell(i, 3)
        c.fill, c.border = ISI, KOTAK
        dv2.add(c)
        s.row_dimensions[i].height = 22
    s.cell(13, 2, "Skala: 1 = sangat tidak setuju, 2 = tidak setuju, 3 = netral, 4 = setuju, 5 = sangat setuju").font = Font(italic=True)
    s.cell(14, 2, "Sumber butir: Sharfina dan Santoso (2016), adaptasi SUS bahasa Indonesia.").font = Font(italic=True, size=9)

    for ws in (p, w, s):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    FORMULIR.mkdir(parents=True, exist_ok=True)
    tpl = FORMULIR / "Formulir Uji Prototipe - TEMPLATE.xlsx"
    wb.save(tpl)
    for i in range(1, JUMLAH + 1):
        tujuan = FORMULIR / f"Formulir Uji Prototipe - EVAL{i:02d}.xlsx"
        if tujuan.exists():
            sys.exit(f"{tujuan.name} sudah ada (mungkin sudah terisi). Tidak ditimpa.")
        shutil.copy(tpl, tujuan)
    print(f"Formulir dibuat di {FORMULIR} (EVAL01-EVAL{JUMLAH:02d}); prototipe acuan: {folder.name}")
    for x in sk:
        print(f"  {x[0]}  {x[5]}")


if __name__ == "__main__":
    main()
