#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
15_analisis_bab4.py
===================
Seluruh analisis kuantitatif Bab IV dari run dan formulir yang SUDAH ADA (tidak ada inferensi
model baru). Mengikuti rancangan evaluasi Bab III (3.2.5) dan daftar kewajiban di dokumen
project `rancangan-penulisan-buku.md` §13.

Keluaran: KOMPONEN_7/analisis_<YYYYmmdd-HHMMSS>/  (folder baru setiap kali dijalankan; tidak
menimpa hasil lama)
  00_ringkasan_analisis.md      ringkasan semua angka utama, siap dibaca room BAB IV
  T1_*.csv                      tujuan 1: kesepakatan penilai + asosiasi fitur-kelas (tau-b)
  T2a_*.csv                     tujuan 2(a): confusion matrix, metrik, bootstrap, McNemar
  T2b_*.csv                     tujuan 2(b): BLEU / METEOR / ROUGE, skor antarmanusia, penyalinan prompt
  T2c_*.csv                     tujuan 2(c): QUEST + Krippendorff alpha
  P_*.csv                       ukuran pendukung: kepatuhan format, waktu komputasi, memori GPU
  G*_*.png                      gambar untuk buku (300 dpi)
  manifest_analisis.json        sidik berkas masukan + versi pustaka (keterulangan)

Prinsip:
  * Run final ditetapkan EKSPLISIT (konstanta RUN di bawah), bukan "terbaru", agar keterulangan terjamin.
  * n analisis sampel besar = 385 (400 dikurangi 15 bangunan set pengembangan; set pengembangan
    diturunkan dengan aturan yang sama persis dengan notebook 09p/09s: benih 42, 5 per kelas xBD dari
    sampel besar di luar sampel kecil).
  * Jawaban kelas 'tidak_jelas' atau gagal urai dihitung SALAH pada metrik utama.
  * Metrik teks: huruf kecil, token = rangkaian huruf/angka (tanpa stemmer, tanpa sinonim), sehingga
    METEOR di sini praktis = METEOR exact match (Banerjee dan Lavie, 2005, modul exact).
  * Semua bootstrap: persentil 95%, 2.000 ulangan, benih 42.

Penggunaan (dari folder KOMPONEN_4):
    pip install pandas numpy scipy scikit-learn matplotlib openpyxl krippendorff
    python 15_analisis_bab4.py
"""

from __future__ import annotations

import datetime
import hashlib
import json
import math
import os
import re
import sys
import warnings
from collections import Counter
from pathlib import Path

try:
    import numpy as np
    import pandas as pd
    import openpyxl
    import krippendorff
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    import scipy
    import sklearn
    from scipy import stats
    from sklearn.metrics import (accuracy_score, balanced_accuracy_score, cohen_kappa_score,
                                 confusion_matrix, f1_score, precision_recall_fscore_support)
except ImportError as e:  # pragma: no cover
    sys.exit(f"Pustaka belum lengkap ({e}). Jalankan:\n"
             "  pip install pandas numpy scipy scikit-learn matplotlib openpyxl krippendorff")

AKAR = Path(os.environ.get("SKRIPSI_AKAR",
            r"F:\KULIAH STIS\TINGKAT IV - SD\SKRIPSWEET\MANUVER SKRIPSI AFTER SENSUS"))
RUNS = AKAR / "KOMPONEN_5" / "runs"
SAMPEL = AKAR / "KOMPONEN_3" / "Output Sampel"
FORMULIR = AKAR / "KOMPONEN_6" / "Formulir"
PENILAIAN_FITUR = AKAR / "KOMPONEN_1"
KELUAR_AKAR = AKAR / "KOMPONEN_7"

# ---- run final (ditetapkan eksplisit) -------------------------------------------------------
RUN = {
    "Qwen2.5-VL-3B": "20260924-145022_qwen3b_besar_v4_b82f43f2",
    "InternVL3-2B": "20260924-213226_ivl2b_besar_v4_b82f43f2",
    "Qwen2.5-VL-3B + LoRA": "20260925-100029_qwen3b-lora_besar_v4_b82f43f2",
    "gpt-6-sol": "20260924-162858_oai-6-sol_kecil_v4_b82f43f2",
}
NAMA_KUNCI_B = {"qwen2.5-vl-3b": "Qwen2.5-VL-3B", "internvl3-2b": "InternVL3-2B", "gpt-6-sol": "gpt-6-sol"}
SVLM_ZS = ["Qwen2.5-VL-3B", "InternVL3-2B"]
MODEL_385 = ["Qwen2.5-VL-3B", "InternVL3-2B", "Qwen2.5-VL-3B + LoRA"]
MODEL_40 = ["Qwen2.5-VL-3B", "InternVL3-2B", "Qwen2.5-VL-3B + LoRA", "gpt-6-sol"]
MODEL_TEKS = ["Qwen2.5-VL-3B", "InternVL3-2B", "gpt-6-sol"]   # LoRA: 0/40 deskripsi sah

KELAS = ["tidak_rusak", "rusak", "hancur"]
PETA_XBD = {"no-damage": "tidak_rusak", "major-damage": "rusak", "destroyed": "hancur"}
FITUR = ["keutuhan_atap", "puing_sekitar", "wujud_bangunan", "sedimen_lumpur"]
KANDIDAT_FORMULIR = ["keutuhan_atap", "puing_sekitar", "wujud_bangunan", "perubahan_radiometri",
                     "perubahan_tekstur", "sedimen_lumpur", "bentuk_dinding", "bangunan_bergeser"]
BENIH, N_BOOT = 42, 2000
warnings.filterwarnings("ignore", category=UserWarning)
warnings.filterwarnings("ignore", category=RuntimeWarning)

# ---- gaya gambar (palet kategoris tervalidasi; urutan tetap, warna mengikuti model) ----------
WARNA = {"Qwen2.5-VL-3B": "#2a78d6", "InternVL3-2B": "#eb6834", "gpt-6-sol": "#1baf7a",
         "Qwen2.5-VL-3B + LoRA": "#eda100"}
TINTA, TINTA2, GARIS = "#0b0b0b", "#52514e", "#d9d8d4"
plt.rcParams.update({
    "font.family": "serif", "font.serif": ["Times New Roman", "Liberation Serif", "DejaVu Serif"],
    "font.size": 10, "axes.edgecolor": TINTA2, "axes.labelcolor": TINTA, "xtick.color": TINTA2,
    "ytick.color": TINTA2, "axes.spines.top": False, "axes.spines.right": False,
    "savefig.dpi": 300, "savefig.bbox": "tight"})


# =============================================================================================
# utilitas
# =============================================================================================
def sha(p: Path, n: int = 12) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()[:n]


def berkas_terurai(d: Path) -> Path:
    calon = sorted(d.glob("hasil_terurai_p*.csv"), key=lambda f: f.stat().st_mtime)
    return calon[-1] if calon else d / "hasil_terurai.csv"


def f3(x, n=3):
    return "–" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:.{n}f}".replace(".", ",")


KOMA = FuncFormatter(lambda v, _: (f"{v:g}").replace(".", ","))


def fp(p):
    return "< 0,0001" if p < 1e-4 else f3(p, 4)


def bootstrap(fungsi, n_unit: int, rng: np.random.Generator) -> tuple[float, float]:
    nilai = []
    for _ in range(N_BOOT):
        i = rng.integers(0, n_unit, n_unit)
        try:
            v = fungsi(i)
        except Exception:
            v = np.nan
        nilai.append(v)
    nilai = np.array(nilai, dtype=float)
    return float(np.nanpercentile(nilai, 2.5)), float(np.nanpercentile(nilai, 97.5))


def mcnemar(benar_a: np.ndarray, benar_b: np.ndarray) -> dict:
    b = int(np.sum(benar_a & ~benar_b))   # hanya A benar
    c = int(np.sum(~benar_a & benar_b))   # hanya B benar
    if b + c == 0:
        return {"hanya_A_benar": b, "hanya_B_benar": c, "uji": "-", "statistik": np.nan, "p": 1.0}
    if b + c < 25:
        p = float(stats.binomtest(b, b + c, 0.5).pvalue)
        return {"hanya_A_benar": b, "hanya_B_benar": c, "uji": "eksak binomial", "statistik": np.nan, "p": p}
    s = (abs(b - c) - 1) ** 2 / (b + c)
    return {"hanya_A_benar": b, "hanya_B_benar": c, "uji": "chi-kuadrat terkoreksi", "statistik": s,
            "p": float(stats.chi2.sf(s, 1))}


def holm(p: list[float]) -> list[float]:
    urut = np.argsort(p)
    m, hasil, maks = len(p), [0.0] * len(p), 0.0
    for k, i in enumerate(urut):
        maks = max(maks, min(1.0, (m - k) * p[i]))
        hasil[i] = maks
    return hasil


# =============================================================================================
# muat data
# =============================================================================================
def muat_sampel():
    besar = pd.read_csv(SAMPEL / "sampel_besar.csv")
    kecil = pd.read_csv(SAMPEL / "sampel_kecil.csv")
    calon = besar[~besar["uid"].isin(kecil["uid"])]
    dev = pd.concat([g.sample(5, random_state=42) for _, g in calon.groupby("kelas_xbd")])
    return besar, kecil, dev


def muat_run(nama: str) -> pd.DataFrame:
    d = RUNS / RUN[nama]
    df = pd.read_csv(berkas_terurai(d)).drop_duplicates("uid")
    df["label"] = df["kelas_xbd"].map(PETA_XBD)
    df["prediksi"] = df["kelas_kerusakan"].where(df["kelas_kerusakan"].isin(KELAS), "salah")
    return df.set_index("uid")


def muat_konfig() -> dict:
    return json.loads((RUNS / RUN["Qwen2.5-VL-3B"] / "paket" / "konfigurasi_fitur.json").read_text(encoding="utf-8"))


# =============================================================================================
# TUJUAN 1
# =============================================================================================
def skor_formulir(f: Path) -> dict:
    ws = openpyxl.load_workbook(f, data_only=True)["Penilaian Fitur"]
    skor = {}
    for r in range(1, ws.max_row + 1):
        no = ws.cell(r, 1).value
        try:
            no = int(float(no))
        except (TypeError, ValueError):
            continue
        if 1 <= no <= 8 and ws.cell(r, 8).value not in (None, ""):
            skor[KANDIDAT_FORMULIR[no - 1]] = float(ws.cell(r, 8).value)
    return skor


def tujuan1(out: Path, runs: dict, konfig: dict, dev_uid: set, L: list):
    L.append("## Tujuan 1 — Kesepakatan penilai kepentingan fitur\n")
    penilai = {"Peneliti": PENILAIAN_FITUR / "Formulir Seleksi Fitur Penting - PENILAIAN PENELITI.xlsx",
               "Ahli BPS MSSD": PENILAIAN_FITUR / "Formulir Seleksi Fitur Penting - PENILAIAN BPS.xlsx",
               "Dosen pembimbing": PENILAIAN_FITUR / "Formulir Seleksi Fitur Penting - PENILAIAN DOSEN.xlsx"}
    skor = {}
    for nama, f in penilai.items():
        if f.exists():
            s = skor_formulir(f)
            if len(s) == 8:
                skor[nama] = s
    label = {c["kunci"]: c["label"] for c in konfig["kandidat"]}
    tab = pd.DataFrame(skor).reindex(KANDIDAT_FORMULIR)
    for p in skor:
        tab[f"peringkat_{p}"] = tab[p].rank(method="min", ascending=False).astype(int)
        tab[f"empat_besar_{p}"] = tab[f"peringkat_{p}"] <= 4
    nama_p = list(skor)
    if len(nama_p) >= 2:
        a, b = nama_p[0], nama_p[1]
        tab["selisih_skor"] = tab[b] - tab[a]
    tab.insert(0, "kandidat", [label.get(k, k) for k in tab.index])
    tab.to_csv(out / "T1_skor_penilai.csv", encoding="utf-8-sig")
    L.append(f"Penilai dengan formulir lengkap: {', '.join(nama_p)} (formulir dosen kosong → tidak dianalisis).\n")
    L.append("| Kandidat | " + " | ".join(nama_p) + " | Selisih |\n|---|" + "---|" * (len(nama_p) + 1))
    for k, r in tab.iterrows():
        L.append(f"| {r['kandidat']} | " + " | ".join(f"{int(r[p])} (#{int(r['peringkat_'+p])})" for p in nama_p)
                 + f" | {int(r.get('selisih_skor', 0)):+d} |")
    ringkas = {}
    if len(nama_p) >= 2:
        x, y = tab[nama_p[0]].values, tab[nama_p[1]].values
        rho, prho = stats.spearmanr(x, y)
        tau, ptau = stats.kendalltau(x, y)
        irisan = set(tab.index[tab[f"empat_besar_{nama_p[0]}"]]) & set(tab.index[tab[f"empat_besar_{nama_p[1]}"]])
        alfa = krippendorff.alpha(reliability_data=np.vstack([x, y]), level_of_measurement="ordinal")
        ringkas = {"n_kandidat": 8, "spearman_rho": rho, "p_rho": prho, "kendall_tau_b": tau, "p_tau": ptau,
                   "krippendorff_alpha_ordinal": alfa, "irisan_empat_besar": f"{len(irisan)} dari 4",
                   "empat_besar_bersama": ", ".join(label[k] for k in KANDIDAT_FORMULIR if k in irisan),
                   **{f"rerata_{p}": float(tab[p].mean()) for p in nama_p}}
        pd.DataFrame([ringkas]).to_csv(out / "T1_kesepakatan.csv", index=False, encoding="utf-8-sig")
        L.append(f"\nIrisan empat besar: **{ringkas['irisan_empat_besar']}** ({ringkas['empat_besar_bersama']}). "
                 f"Spearman ρ = {f3(rho)} (p = {f3(prho, 4)}); Kendall τ-b = {f3(tau)} (p = {f3(ptau, 4)}); "
                 f"Krippendorff α ordinal = {f3(alfa)}. Rerata skor: "
                 + "; ".join(f"{p} {f3(tab[p].mean(), 2)}" for p in nama_p) + ". n = 8 butir → koefisien pelengkap, "
                 "bukan dasar keputusan.\n")

    # ---- asosiasi fitur-kelas (validasi empiris) ----
    baris = []
    for nama in ["Qwen2.5-VL-3B", "InternVL3-2B", "gpt-6-sol"]:
        df = runs[nama]
        df = df[~df.index.isin(dev_uid)] if nama != "gpt-6-sol" else df
        lvl = {c["kunci"]: c["level"] for c in konfig["kandidat"] if c["kunci"] in FITUR}
        y_xbd = df["label"].map({k: i for i, k in enumerate(KELAS)})
        y_mod = df["kelas_kerusakan"].map({k: i for i, k in enumerate(KELAS)})
        for f in FITUR:
            x = df[f"fitur_{f}"].map({v: i for i, v in enumerate(lvl[f])})
            for acuan, y in (("label xBD", y_xbd), ("kelas model", y_mod)):
                m = x.notna() & y.notna()
                n = int(m.sum())
                sebaran = df.loc[m, f"fitur_{f}"].value_counts().to_dict()
                if n < 3 or x[m].nunique() < 2 or y[m].nunique() < 2:
                    tau, p, ket = np.nan, np.nan, "tidak dapat dihitung (jawaban fitur atau acuan konstan)"
                else:
                    tau, p = stats.kendalltau(x[m], y[m])
                    ket = ""
                baris.append({"model": nama, "fitur": f, "acuan": acuan, "n_sah": n,
                              "tidak_jelas_atau_kosong": int(len(df) - x.notna().sum()),
                              "tau_b": tau, "p": p, "sebaran_jawaban": json.dumps(sebaran, ensure_ascii=False),
                              "keterangan": ket})
    ta = pd.DataFrame(baris)
    ta.to_csv(out / "T1_asosiasi_fitur_kelas.csv", index=False, encoding="utf-8-sig")
    L.append("\n### Asosiasi fitur–kelas (Kendall τ-b; fitur ordinal keluaran model vs label xBD)\n")
    L.append("| Model | n | " + " | ".join(FITUR) + " |\n|---|---|" + "---|" * len(FITUR))
    for nama in ["Qwen2.5-VL-3B", "InternVL3-2B", "gpt-6-sol"]:
        s = ta[(ta.model == nama) & (ta.acuan == "label xBD")].set_index("fitur")
        L.append(f"| {nama} | {int(s['n_sah'].max())} | " + " | ".join(
            (f"{f3(s.loc[f,'tau_b'])}{'*' if s.loc[f,'p'] < 0.05 else ''}" if not math.isnan(s.loc[f, 'tau_b'])
             else "t.d.") for f in FITUR) + " |")
    L.append("\n`*` p < 0,05; t.d. = tidak dapat dihitung karena jawaban fitur konstan. LoRA tidak diikutkan "
             "(0/40 keluaran fitur sah). Rincian + τ-b terhadap kelas model sendiri: `T1_asosiasi_fitur_kelas.csv`.\n")


# =============================================================================================
# TUJUAN 2 (a) — confusion matrix
# =============================================================================================
def metrik(y: np.ndarray, p: np.ndarray) -> dict:
    lab_semua = KELAS + ["salah"]
    pr, rc, f1, _ = precision_recall_fscore_support(y, p, labels=KELAS, zero_division=0)
    valid = np.isin(p, KELAS)
    num = {k: i for i, k in enumerate(KELAS)}
    kw = (cohen_kappa_score([num[v] for v in y[valid]], [num[v] for v in p[valid]], weights="linear")
          if valid.sum() > 1 else np.nan)
    yb, pb = (y != "tidak_rusak"), np.isin(p, ["rusak", "hancur"])
    return {"n": len(y), "n_tidak_sah": int((~valid).sum()),
            "akurasi": accuracy_score(y, p), "akurasi_seimbang": balanced_accuracy_score(y, p),
            "f1_makro": f1_score(y, p, labels=KELAS, average="macro", zero_division=0),
            "kappa": cohen_kappa_score(y, p, labels=lab_semua), "kappa_linear_sah": kw,
            **{f"presisi_{k}": pr[i] for i, k in enumerate(KELAS)},
            **{f"recall_{k}": rc[i] for i, k in enumerate(KELAS)},
            **{f"f1_{k}": f1[i] for i, k in enumerate(KELAS)},
            "biner_akurasi": accuracy_score(yb, pb), "biner_akurasi_seimbang": balanced_accuracy_score(yb, pb),
            "biner_f1_rusak_atau_hancur": f1_score(yb, pb, zero_division=0),
            "biner_kappa": cohen_kappa_score(yb, pb)}


def tujuan2a(out: Path, runs: dict, uid385: list, uid40: list, L: list):
    L.append("\n## Tujuan 2(a) — Confusion matrix terhadap label xBD\n")
    rng = np.random.default_rng(BENIH)
    baris, cm_simpan = [], {}
    for himp, uid, model in (("n385", uid385, MODEL_385), ("n40", uid40, MODEL_40)):
        for nama in model:
            df = runs[nama].loc[uid]
            y, p = df["label"].values, df["prediksi"].values
            m = metrik(y, p)
            for kunci, fn in (("akurasi", accuracy_score), ("akurasi_seimbang", balanced_accuracy_score),
                              ("f1_makro", lambda a, b: f1_score(a, b, labels=KELAS, average="macro", zero_division=0)),
                              ("kappa", lambda a, b: cohen_kappa_score(a, b, labels=KELAS + ["salah"]))):
                lo, hi = bootstrap(lambda i: fn(y[i], p[i]), len(y), rng)
                m[f"{kunci}_ci_bawah"], m[f"{kunci}_ci_atas"] = lo, hi
            m["sebaran_prediksi"] = json.dumps(pd.Series(p).value_counts().to_dict(), ensure_ascii=False)
            baris.append({"himpunan": himp, "model": nama, **m})
            cm = pd.DataFrame(confusion_matrix(y, p, labels=KELAS + ["salah"]),
                              index=[f"xBD:{k}" for k in KELAS + ["salah"]], columns=KELAS + ["tidak_sah"]).iloc[:3]
            cm.to_csv(out / f"T2a_matriks_{himp}_{nama.replace(' ', '').replace('+', '_')}.csv", encoding="utf-8-sig")
            cm_simpan[(himp, nama)] = cm
    tm = pd.DataFrame(baris)
    tm.to_csv(out / "T2a_metrik_klasifikasi.csv", index=False, encoding="utf-8-sig")

    for himp in ("n385", "n40"):
        s = tm[tm.himpunan == himp]
        L.append(f"\n**{'Sampel besar tanpa set pengembangan (n = 385)' if himp == 'n385' else 'Sampel kecil (n = 40)'}**\n")
        L.append("| Model | Akurasi | Akurasi seimbang [95%] | F1 makro [95%] | κ | κ linear | F1 tr/r/h | Biner F1 | tidak sah |\n"
                 "|---|---|---|---|---|---|---|---|---|")
        for _, r in s.iterrows():
            L.append(f"| {r.model} | {f3(r.akurasi)} | {f3(r.akurasi_seimbang)} [{f3(r.akurasi_seimbang_ci_bawah)}–"
                     f"{f3(r.akurasi_seimbang_ci_atas)}] | {f3(r.f1_makro)} [{f3(r.f1_makro_ci_bawah)}–{f3(r.f1_makro_ci_atas)}]"
                     f" | {f3(r.kappa)} | {f3(r.kappa_linear_sah)} | {f3(r.f1_tidak_rusak,2)}/{f3(r.f1_rusak,2)}/"
                     f"{f3(r.f1_hancur,2)} | {f3(r.biner_f1_rusak_atau_hancur)} | {r.n_tidak_sah} |")
    L.append("\nκ memakai 4 kategori (tidak sah = kategori tersendiri); κ linear hanya pada jawaban sah. "
             "Pembanding acak untuk akurasi seimbang 3 kelas = 1/3.\n")

    # McNemar
    uji = []
    pasangan = [("n385", "Qwen2.5-VL-3B + LoRA", "Qwen2.5-VL-3B"), ("n385", "Qwen2.5-VL-3B + LoRA", "InternVL3-2B"),
                ("n385", "InternVL3-2B", "Qwen2.5-VL-3B"), ("n40", "gpt-6-sol", "Qwen2.5-VL-3B"),
                ("n40", "gpt-6-sol", "InternVL3-2B"), ("n40", "Qwen2.5-VL-3B + LoRA", "gpt-6-sol")]
    for himp, a, b in pasangan:
        uid = uid385 if himp == "n385" else uid40
        da, db = runs[a].loc[uid], runs[b].loc[uid]
        r = mcnemar((da.prediksi == da.label).values, (db.prediksi == db.label).values)
        uji.append({"himpunan": himp, "model_A": a, "model_B": b, **r})
    tu = pd.DataFrame(uji)
    tu.to_csv(out / "T2a_mcnemar.csv", index=False, encoding="utf-8-sig")
    L.append("**Uji McNemar (α = 5%)**\n\n| Himpunan | A vs B | hanya A benar | hanya B benar | uji | p |\n|---|---|---|---|---|---|")
    for _, r in tu.iterrows():
        L.append(f"| {r.himpunan} | {r.model_A} vs {r.model_B} | {r.hanya_A_benar} | {r.hanya_B_benar} | {r.uji} | "
                 f"{fp(r.p)} |")

    # gambar confusion matrix
    fig, ax = plt.subplots(1, 3, figsize=(10.5, 3.4), gridspec_kw={"wspace": 0.08})
    for k_ax, (a, nama) in enumerate(zip(ax, MODEL_385)):
        cm = cm_simpan[("n385", nama)][KELAS].values
        prop = cm / cm.sum(1, keepdims=True)
        a.imshow(prop, cmap="Blues", vmin=0, vmax=1)
        for i in range(3):
            for j in range(3):
                a.text(j, i, f"{cm[i, j]}\n({prop[i, j]*100:.0f}%)".replace(".", ","), ha="center", va="center",
                       fontsize=8.5, color="white" if prop[i, j] > 0.55 else TINTA)
        a.set_xticks(range(3), ["tidak rusak", "rusak", "hancur"], fontsize=8.5)
        a.set_yticks(range(3), ["no-damage", "major-damage", "destroyed"] if k_ax == 0 else [""] * 3, fontsize=8.5)
        a.tick_params(length=0)
        a.set_xlabel("Kelas keluaran model")
        a.set_title(nama, fontsize=10, color=TINTA)
        for s in a.spines.values():
            s.set_visible(False)
    ax[0].set_ylabel("Label xBD")
    fig.savefig(out / "G1_confusion_matrix_n385.png")
    plt.close(fig)

    # gambar metrik ringkas
    s = tm[tm.himpunan == "n385"].set_index("model")
    fig, ax = plt.subplots(figsize=(6.6, 3.2))
    met = [("akurasi_seimbang", "Akurasi seimbang"), ("f1_makro", "F1 makro")]
    lebar = 0.26
    for j, nama in enumerate(MODEL_385):
        for i, (k, _) in enumerate(met):
            v = s.loc[nama, k]
            x = i + (j - 1) * (lebar + 0.02)
            ax.bar(x, v, lebar, color=WARNA[nama], label=nama if i == 0 else None)
            ax.errorbar(x, v, yerr=[[v - s.loc[nama, f"{k}_ci_bawah"]], [s.loc[nama, f"{k}_ci_atas"] - v]],
                        color=TINTA2, capsize=2, lw=0.8)
            ax.text(x, s.loc[nama, f"{k}_ci_atas"] + 0.015, f3(v, 2), ha="center", fontsize=8, color=TINTA)
    ax.axhline(1 / 3, color=TINTA2, lw=0.8, ls="--")
    ax.text(1.55, 1 / 3 + 0.01, "acak = 1/3", fontsize=8, color=TINTA2, ha="right")
    ax.set_xticks(range(len(met)), [m[1] for m in met])
    ax.set_ylim(0, 0.9)
    ax.yaxis.set_major_formatter(KOMA)
    ax.yaxis.grid(True, color=GARIS, lw=0.6)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=8.5, loc="upper left")
    fig.savefig(out / "G2_metrik_kelas_n385.png")
    plt.close(fig)
    return tm


# =============================================================================================
# TUJUAN 2 (b) — metrik teks
# =============================================================================================
def token(t) -> list[str]:
    return re.findall(r"[0-9a-zà-ÿ]+", str(t).lower()) if isinstance(t, str) and t.strip() else []


def ngram(tok, n):
    return Counter(tuple(tok[i:i + n]) for i in range(len(tok) - n + 1))


def bleu_korpus(hip: list[list[str]], refs: list[list[list[str]]], N: int) -> float:
    cocok, total = [0] * N, [0] * N
    c_len, r_len = 0, 0
    for h, rs in zip(hip, refs):
        c_len += len(h)
        r_len += min((abs(len(r) - len(h)), len(r)) for r in rs)[1] if rs else 0
        for n in range(1, N + 1):
            hn = ngram(h, n)
            maks = Counter()
            for r in rs:
                for g, k in ngram(r, n).items():
                    maks[g] = max(maks[g], k)
            cocok[n - 1] += sum(min(k, maks[g]) for g, k in hn.items())
            total[n - 1] += max(len(h) - n + 1, 0)
    if c_len == 0 or any(t == 0 for t in total) or any(c == 0 for c in cocok):
        return 0.0
    bp = 1.0 if c_len > r_len else math.exp(1 - r_len / c_len)
    return bp * math.exp(sum(math.log(c / t) for c, t in zip(cocok, total)) / N)


def meteor_exact(h: list[str], r: list[str]) -> float:
    if not h or not r:
        return 0.0
    sisa, jajar = list(range(len(r))), []
    for i, w in enumerate(h):                      # penjajaran exact, kiri ke kanan
        for j in sisa:
            if r[j] == w:
                jajar.append((i, j)); sisa.remove(j); break
    m = len(jajar)
    if m == 0:
        return 0.0
    P, R = m / len(h), m / len(r)
    fmean = 10 * P * R / (R + 9 * P)
    chunk = 1
    for (i1, j1), (i2, j2) in zip(jajar, jajar[1:]):
        if not (i2 == i1 + 1 and j2 == j1 + 1):
            chunk += 1
    return fmean * (1 - 0.5 * (chunk / m) ** 3)


def rouge_n(h, r, n) -> tuple[float, float, float]:
    hn, rn = ngram(h, n), ngram(r, n)
    over = sum((hn & rn).values())
    P = over / max(sum(hn.values()), 1)
    R = over / max(sum(rn.values()), 1)
    return P, R, (2 * P * R / (P + R) if P + R else 0.0)


def rouge_l(h, r) -> tuple[float, float, float]:
    if not h or not r:
        return 0.0, 0.0, 0.0
    t = [[0] * (len(r) + 1) for _ in range(len(h) + 1)]
    for i in range(len(h)):
        for j in range(len(r)):
            t[i + 1][j + 1] = t[i][j] + 1 if h[i] == r[j] else max(t[i][j + 1], t[i + 1][j])
    lcs = t[-1][-1]
    P, R = lcs / len(h), lcs / len(r)
    return P, R, (2 * P * R / (P + R) if P + R else 0.0)


def skor_teks(hip: list[list[str]], refs: list[list[list[str]]]) -> dict:
    s = {"BLEU-1": bleu_korpus(hip, refs, 1), "BLEU-4": bleu_korpus(hip, refs, 4)}
    per = {"METEOR": [], "ROUGE-1": [], "ROUGE-2": [], "ROUGE-L": [], "ROUGE-1_R": [], "ROUGE-L_R": []}
    for h, rs in zip(hip, refs):
        per["METEOR"].append(max(meteor_exact(h, r) for r in rs))
        r1 = max((rouge_n(h, r, 1) for r in rs), key=lambda x: x[2])
        r2 = max((rouge_n(h, r, 2) for r in rs), key=lambda x: x[2])
        rl = max((rouge_l(h, r) for r in rs), key=lambda x: x[2])
        per["ROUGE-1"].append(r1[2]); per["ROUGE-2"].append(r2[2]); per["ROUGE-L"].append(rl[2])
        per["ROUGE-1_R"].append(r1[1]); per["ROUGE-L_R"].append(rl[1])
    s.update({k: float(np.mean(v)) for k, v in per.items()})
    return s


def baca_sesi_a() -> pd.DataFrame:
    kunci = pd.read_csv(FORMULIR / "kunci_formulir_sesi_a.csv")
    baris = []
    for i in range(1, 6):
        f = FORMULIR / f"Formulir Sesi A - EVAL{i:02d}.xlsx"
        ws = openpyxl.load_workbook(f, read_only=True)["Deskripsi"]
        for r in range(2, 42):
            baris.append({"evaluator": f"EVAL{i:02d}", "kode": ws.cell(r, 2).value,
                          "deskripsi": str(ws.cell(r, 4).value or "").strip()})
    return pd.DataFrame(baris).merge(kunci[["kode", "uid", "kelas_xbd"]], on="kode")


def tujuan2b(out: Path, runs: dict, uid40: list, konfig: dict, L: list):
    L.append("\n## Tujuan 2(b) — Evaluasi otomatis deskripsi (acuan: 5 deskripsi Sesi A)\n")
    sa = baca_sesi_a()
    sa.to_csv(out / "T2b_deskripsi_referensi_sesi_a.csv", index=False, encoding="utf-8-sig")
    ev = sorted(sa.evaluator.unique())
    ref = {u: {e: token(t) for e, t in zip(g.evaluator, g.deskripsi)} for u, g in sa.groupby("uid")}
    uid = [u for u in uid40 if u in ref]

    baris = []
    for nama in MODEL_TEKS:
        df = runs[nama].loc[uid]
        hip = [token(t) for t in df["deskripsi"]]
        kosong = sum(1 for h in hip if not h)
        s5 = skor_teks(hip, [list(ref[u].values()) for u in uid])
        baris.append({"sistem": nama, "protokol": "vs 5 acuan", "n": len(uid), "deskripsi_kosong": kosong, **s5})
        # protokol setara manusia: rerata atas 5 kali 'tinggalkan satu evaluator'
        acc = [skor_teks(hip, [[ref[u][e2] for e2 in ev if e2 != e] for u in uid]) for e in ev]
        baris.append({"sistem": nama, "protokol": "vs 4 acuan (rerata 5 lipatan)", "n": len(uid),
                      "deskripsi_kosong": kosong, **{k: float(np.mean([a[k] for a in acc])) for k in acc[0]}})
    # antarmanusia
    acc = []
    for e in ev:
        hip = [ref[u][e] for u in uid]
        acc.append(skor_teks(hip, [[ref[u][e2] for e2 in ev if e2 != e] for u in uid]))
        baris.append({"sistem": f"Manusia {e}", "protokol": "vs 4 acuan", "n": len(uid), "deskripsi_kosong": 0, **acc[-1]})
    baris.append({"sistem": "Manusia (rerata)", "protokol": "vs 4 acuan", "n": len(uid), "deskripsi_kosong": 0,
                  **{k: float(np.mean([a[k] for a in acc])) for k in acc[0]}})
    tt = pd.DataFrame(baris)
    tt.to_csv(out / "T2b_metrik_teks.csv", index=False, encoding="utf-8-sig")

    # keragaman & penyalinan prompt
    prompt = (RUNS / RUN["Qwen2.5-VL-3B"] / "paket" / "prompt_fitur_v4_b82f43f2.txt").read_text(encoding="utf-8")
    definisi = " ".join(c["definisi"] for c in konfig["kandidat"] if c["kunci"] in FITUR)
    g5_prompt = set(ngram(token(prompt), 5)) | set(ngram(token(definisi), 5))
    g4_prompt = set(ngram(token(prompt), 4)) | set(ngram(token(definisi), 4))

    def salin(tok):
        if len(tok) < 5:
            return False, 0.0
        g4 = list(ngram(tok, 4).keys())
        return bool(set(ngram(tok, 5)) & g5_prompt), (sum(g in g4_prompt for g in g4) / len(g4) if g4 else 0.0)

    kr = []
    for nama in MODEL_TEKS:
        toks = [token(t) for t in runs[nama].loc[uid, "deskripsi"]]
        sal = [salin(t) for t in toks]
        kr.append({"sistem": nama, "kalimat_unik": len({" ".join(t) for t in toks}),
                   "rerata_kata": float(np.mean([len(t) for t in toks])),
                   "proporsi_memuat_5gram_prompt": float(np.mean([s[0] for s in sal])),
                   "rerata_proporsi_4gram_dari_prompt": float(np.mean([s[1] for s in sal]))})
    for e in ev:
        toks = [ref[u][e] for u in uid]
        sal = [salin(t) for t in toks]
        kr.append({"sistem": f"Manusia {e}", "kalimat_unik": len({" ".join(t) for t in toks}),
                   "rerata_kata": float(np.mean([len(t) for t in toks])),
                   "proporsi_memuat_5gram_prompt": float(np.mean([s[0] for s in sal])),
                   "rerata_proporsi_4gram_dari_prompt": float(np.mean([s[1] for s in sal]))})
    tk = pd.DataFrame(kr)
    tk.to_csv(out / "T2b_keragaman_penyalinan.csv", index=False, encoding="utf-8-sig")

    kol = ["BLEU-1", "BLEU-4", "METEOR", "ROUGE-1", "ROUGE-2", "ROUGE-L"]
    L.append("Huruf kecil, tanpa stemmer/sinonim (METEOR = exact match), ROUGE = F1 terbaik antaracuan; BLEU korpus.\n")
    L.append("| Sistem | Protokol | " + " | ".join(kol) + " |\n|---|---|" + "---|" * len(kol))
    for _, r in tt[~tt.sistem.str.match(r"Manusia EVAL")].iterrows():
        L.append(f"| {r.sistem} | {r.protokol} | " + " | ".join(f3(r[k]) for k in kol) + " |")
    L.append("\nPerbandingan setara: baris 'vs 4 acuan' (model) dibandingkan dengan 'Manusia (rerata)'. Rincian per "
             "evaluator: `T2b_metrik_teks.csv`.\n")
    L.append("| Sistem | Kalimat unik /40 | Rerata kata | Memuat 5-gram prompt | Proporsi 4-gram dari prompt |\n|---|---|---|---|---|")
    for _, r in tk.iterrows():
        L.append(f"| {r.sistem} | {r.kalimat_unik} | {f3(r.rerata_kata,1)} | {f3(r.proporsi_memuat_5gram_prompt,2)} | "
                 f"{f3(r.rerata_proporsi_4gram_dari_prompt,2)} |")

    # gambar
    sub = tt[(tt.protokol != "vs 5 acuan") & (~tt.sistem.str.match(r"Manusia EVAL"))].set_index("sistem")
    urut = MODEL_TEKS + ["Manusia (rerata)"]
    fig, ax = plt.subplots(figsize=(7.2, 3.2))
    lebar = 0.2
    for j, s in enumerate(urut):
        for i, k in enumerate(["BLEU-4", "METEOR", "ROUGE-1", "ROUGE-L"]):
            v = sub.loc[s, k]
            x = i + (j - 1.5) * (lebar + 0.01)
            ax.bar(x, v, lebar, color=WARNA.get(s, "#9a9994"), label=s if i == 0 else None)
            ax.text(x, v + 0.01, f3(v, 2), ha="center", fontsize=7, color=TINTA)
    ax.set_xticks(range(4), ["BLEU-4", "METEOR", "ROUGE-1", "ROUGE-L"])
    ax.yaxis.set_major_formatter(KOMA)
    ax.yaxis.grid(True, color=GARIS, lw=0.6)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=8, ncol=4, loc="upper center", bbox_to_anchor=(0.5, 1.15))
    fig.savefig(out / "G3_metrik_teks.png")
    plt.close(fig)
    return tt, tk


# =============================================================================================
# TUJUAN 2 (c) — QUEST
# =============================================================================================
def tujuan2c(out: Path, L: list):
    L.append("\n## Tujuan 2(c) — Evaluasi manusia QUEST (Sesi B)\n")
    kunci = pd.read_csv(FORMULIR / "kunci_formulir_sesi_b.csv")
    baris = []
    for i in range(1, 6):
        ws = openpyxl.load_workbook(FORMULIR / f"Formulir Sesi B - EVAL{i:02d}.xlsx", read_only=True)["Penilaian"]
        for r in range(2, 2 + len(kunci)):
            for c, kr in zip((6, 7, 8), ("Akurasi", "Relevansi", "Kepuasan")):
                v = ws.cell(r, c).value
                baris.append({"evaluator": f"EVAL{i:02d}", "baris_excel": r, "kriteria": kr,
                              "skor": v if isinstance(v, (int, float)) and 1 <= v <= 5 else np.nan})
    d = pd.DataFrame(baris).merge(kunci[["baris_excel", "kode", "uid", "kelas_xbd", "model"]], on="baris_excel")
    d["model"] = d["model"].map(NAMA_KUNCI_B)
    d.to_csv(out / "T2c_skor_mentah_panjang.csv", index=False, encoding="utf-8-sig")
    rng = np.random.default_rng(BENIH)
    KR = ["Akurasi", "Relevansi", "Kepuasan"]
    ring = []
    for nama in MODEL_TEKS:
        for kr in KR:
            s = d[(d.model == nama) & (d.kriteria == kr)]
            per_bgn = s.groupby("uid").skor.mean().values
            lo, hi = bootstrap(lambda i: float(np.mean(per_bgn[i])), len(per_bgn), rng)
            ring.append({"model": nama, "kriteria": kr, "rerata": s.skor.mean(), "sd": s.skor.std(),
                         "median": s.skor.median(), "ci_bawah": lo, "ci_atas": hi,
                         "n_skor_sah": int(s.skor.notna().sum()), "n_kosong": int(s.skor.isna().sum())})
    tr = pd.DataFrame(ring)
    tr.to_csv(out / "T2c_rerata_quest.csv", index=False, encoding="utf-8-sig")
    d.pivot_table(index=["model", "kelas_xbd"], columns="kriteria", values="skor", aggfunc="mean").to_csv(
        out / "T2c_rerata_per_kelas_xbd.csv", encoding="utf-8-sig")
    d.pivot_table(index="evaluator", columns=["model"], values="skor", aggfunc="mean").to_csv(
        out / "T2c_rerata_per_evaluator.csv", encoding="utf-8-sig")

    al = []
    for kr in KR:
        m = d[d.kriteria == kr].pivot_table(index="evaluator", columns=["uid", "model"], values="skor").values
        al.append({"kriteria": kr, "alpha_ordinal": krippendorff.alpha(reliability_data=m, level_of_measurement="ordinal"),
                   "alpha_interval": krippendorff.alpha(reliability_data=m, level_of_measurement="interval"),
                   "n_unit": m.shape[1], "n_penilai": m.shape[0]})
    ta = pd.DataFrame(al)
    ta.to_csv(out / "T2c_krippendorff.csv", index=False, encoding="utf-8-sig")

    # pelengkap: Friedman (blok = bangunan, skor = rerata 5 penilai) + Wilcoxon berpasangan (Holm)
    uji = []
    for kr in KR:
        w = d[d.kriteria == kr].pivot_table(index="uid", columns="model", values="skor", aggfunc="mean")[MODEL_TEKS].dropna()
        fr = stats.friedmanchisquare(*[w[m].values for m in MODEL_TEKS])
        pas = [(a, b) for i, a in enumerate(MODEL_TEKS) for b in MODEL_TEKS[i + 1:]]
        pw = [stats.wilcoxon(w[a], w[b]).pvalue for a, b in pas]
        ph = holm(pw)
        for (a, b), p0, p1 in zip(pas, pw, ph):
            uji.append({"kriteria": kr, "friedman_chi2": fr.statistic, "friedman_p": fr.pvalue, "n_bangunan": len(w),
                        "pasangan": f"{a} vs {b}", "wilcoxon_p": p0, "wilcoxon_p_holm": p1})
    tu = pd.DataFrame(uji)
    tu.to_csv(out / "T2c_uji_perbedaan_pelengkap.csv", index=False, encoding="utf-8-sig")

    L.append("| Model | Akurasi [95%] | Relevansi [95%] | Kepuasan [95%] |\n|---|---|---|---|")
    for nama in MODEL_TEKS:
        s = tr[tr.model == nama].set_index("kriteria")
        L.append(f"| {nama} | " + " | ".join(f"{f3(s.loc[k,'rerata'],2)} [{f3(s.loc[k,'ci_bawah'],2)}–{f3(s.loc[k,'ci_atas'],2)}]"
                                           for k in KR) + " |")
    L.append("\nKrippendorff α (5 penilai × 120 unit): " + "; ".join(
        f"{r.kriteria} ordinal {f3(r.alpha_ordinal)} (interval {f3(r.alpha_interval)})" for _, r in ta.iterrows())
             + ". Ambang: ≥ 0,800 andal; 0,667–0,800 tentatif; < 0,667 tidak memadai.\n")
    L.append("Pelengkap (tidak direncanakan di Bab III; boleh dipakai sebagai pendukung): Friedman per kriteria, blok = bangunan: "
             + "; ".join(f"{k} χ² = {f3(g.friedman_chi2.iloc[0],2)}, p {fp(g.friedman_p.iloc[0]) if fp(g.friedman_p.iloc[0]).startswith('<') else '= ' + fp(g.friedman_p.iloc[0])}" for k, g in tu.groupby("kriteria", sort=False))
             + ". Wilcoxon berpasangan (Holm) di `T2c_uji_perbedaan_pelengkap.csv`.\n")

    fig, ax = plt.subplots(figsize=(6.6, 3.2))
    lebar = 0.26
    for j, nama in enumerate(MODEL_TEKS):
        s = tr[tr.model == nama].set_index("kriteria")
        for i, k in enumerate(KR):
            v = s.loc[k, "rerata"]
            x = i + (j - 1) * (lebar + 0.02)
            ax.bar(x, v, lebar, color=WARNA[nama], label=nama if i == 0 else None)
            ax.errorbar(x, v, yerr=[[v - s.loc[k, "ci_bawah"]], [s.loc[k, "ci_atas"] - v]], color=TINTA2, capsize=2, lw=0.8)
            ax.text(x, s.loc[k, "ci_atas"] + 0.06, f3(v, 2), ha="center", fontsize=8, color=TINTA)
    ax.set_xticks(range(3), KR)
    ax.set_ylim(1, 5.3)
    ax.set_ylabel("Rerata skor (1–5)")
    ax.yaxis.grid(True, color=GARIS, lw=0.6)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=8.5, ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.14))
    fig.savefig(out / "G4_quest.png")
    plt.close(fig)
    return tr, ta


# =============================================================================================
# ukuran pendukung
# =============================================================================================
def pendukung(out: Path, runs: dict, L: list):
    L.append("\n## Ukuran pendukung — kepatuhan format dan komputasi\n")
    baris = []
    for nama, run in RUN.items():
        df = runs[nama]
        w = json.loads((RUNS / run / "waktu_komputasi.json").read_text(encoding="utf-8"))
        ada_fitur = df["f_panjang_mentah"].fillna(0) > 0
        f = df[ada_fitur]
        sel_fitur = f[[f"fitur_{x}" for x in FITUR]]
        baris.append({
            "model": nama, "n_panggilan_kelas": len(df), "n_panggilan_fitur": int(ada_fitur.sum()),
            "kelas_terurai": df["k_terurai"].mean(),
            "kelas_tingkat_urai": json.dumps(df["k_tingkat_perbaikan"].fillna("kosong").value_counts().to_dict()),
            "kelas_tidak_jelas": (df["kelas_kerusakan"] == "tidak_jelas").mean(),
            "kelas_tidak_sah": (~df["kelas_kerusakan"].isin(KELAS)).mean(),
            "fitur_terurai": f["f_terurai"].mean() if len(f) else np.nan,
            "fitur_tingkat_urai": json.dumps(f["f_tingkat_perbaikan"].fillna("kosong").value_counts().to_dict()),
            "fitur_lengkap_sah": f["f_sah"].mean() if len(f) else np.nan,
            "sel_fitur_tidak_jelas_atau_kosong": float(sel_fitur.isna().mean().mean() + (sel_fitur == "tidak_jelas").mean().mean()) if len(f) else np.nan,
            "detik_muat_model": w.get("detik_muat_model"),
            "detik_per_panggilan_fitur": (w.get("per_panggilan_fitur") or {}).get("rerata"),
            "detik_per_panggilan_kelas": (w.get("per_panggilan_kelas") or {}).get("rerata"),
            "detik_per_bangunan": (w.get("per_bangunan") or {}).get("rerata"),
            "memori_gpu_puncak_gb": w.get("memori_gpu_puncak_inferensi_gb", w.get("memori_gpu_puncak_uji_gb")),
            "jam_dinding_menit": w.get("jam_dinding_loop_menit", w.get("jam_dinding_menit", w.get("jam_dinding_uji_menit"))),
            "detik_pelatihan": w.get("detik_pelatihan"), "cakupan_waktu": w.get("cakupan", "")})
    tp = pd.DataFrame(baris)
    tp.to_csv(out / "P_kepatuhan_format_komputasi.csv", index=False, encoding="utf-8-sig")
    L.append("| Model | Kelas terurai | Tingkat urai kelas | Kelas tidak sah | Fitur lengkap sah | Tingkat urai fitur | s/panggilan kelas | s/bangunan | GPU puncak (GB) |\n|---|---|---|---|---|---|---|---|---|")
    for _, r in tp.iterrows():
        L.append(f"| {r.model} | {f3(r.kelas_terurai,3)} | {r.kelas_tingkat_urai} | {f3(r.kelas_tidak_sah,3)} | "
                 f"{f3(r.fitur_lengkap_sah,3)} (n={r.n_panggilan_fitur}) | {r.fitur_tingkat_urai} | {f3(r.detik_per_panggilan_kelas,2)} | "
                 f"{f3(r.detik_per_bangunan,2)} | {f3(r.memori_gpu_puncak_gb,2) if r.memori_gpu_puncak_gb else '–'} |")
    L.append("\nTingkat urai: 'langsung' = JSON terbaca apa adanya; 'pagar_kode' = JSON dibungkus blok kode markdown lalu dilepas "
             "pengurai (bukan kesalahan isi); 'gagal' = tidak ada objek JSON yang dapat diurai.")
    L.append("\ngpt-6-sol: latensi API (jaringan + antrean + server), tidak setara dengan waktu GPU lokal. LoRA: waktu "
             f"pelatihan {f3(tp.set_index('model').loc['Qwen2.5-VL-3B + LoRA','detik_pelatihan']/60,1)} menit (T4).\n")


# =============================================================================================
def main():
    for p in (SAMPEL / "sampel_besar.csv", FORMULIR / "kunci_formulir_sesi_a.csv", FORMULIR / "kunci_formulir_sesi_b.csv"):
        if not p.exists():
            sys.exit(f"Berkas tidak ditemukan: {p}")
    waktu = datetime.datetime.now()
    out = KELUAR_AKAR / f"analisis_{waktu:%Y%m%d-%H%M%S}"
    out.mkdir(parents=True)

    besar, kecil, dev = muat_sampel()
    dev_uid = set(dev["uid"])
    runs = {n: muat_run(n) for n in RUN}
    konfig = muat_konfig()
    uid385 = [u for u in besar["uid"] if u not in dev_uid]
    uid40 = list(kecil["uid"])
    for n in MODEL_385:
        hilang = set(uid385) - set(runs[n].index)
        if hilang:
            sys.exit(f"{n}: {len(hilang)} bangunan n=385 tidak ada di run")
    assert len(uid385) == 385 and len(uid40) == 40, (len(uid385), len(uid40))

    L = [f"# Ringkasan Analisis Bab IV\n\nDijalankan {waktu:%d-%m-%Y %H:%M}. Skrip `15_analisis_bab4.py`. "
         f"Angka dengan koma desimal. n = 385 (sampel besar tanpa 15 bangunan set pengembangan) dan n = 40 (sampel kecil).\n"]
    tujuan1(out, runs, konfig, dev_uid, L)
    tujuan2a(out, runs, uid385, uid40, L)
    tujuan2b(out, runs, uid40, konfig, L)
    tujuan2c(out, L)
    pendukung(out, runs, L)
    (out / "00_ringkasan_analisis.md").write_text("\n".join(L) + "\n", encoding="utf-8")

    masukan = [SAMPEL / "sampel_besar.csv", SAMPEL / "sampel_kecil.csv", FORMULIR / "kunci_formulir_sesi_a.csv",
               FORMULIR / "kunci_formulir_sesi_b.csv"] + [berkas_terurai(RUNS / r) for r in RUN.values()] + \
              [FORMULIR / f"Formulir Sesi {s} - EVAL{i:02d}.xlsx" for s in "AB" for i in range(1, 6)]
    (out / "manifest_analisis.json").write_text(json.dumps({
        "waktu": waktu.isoformat(timespec="seconds"), "skrip_sha": sha(Path(__file__)), "run": RUN,
        "benih": BENIH, "n_bootstrap": N_BOOT, "n385": 385, "n40": 40,
        "uid_set_pengembangan": sorted(dev_uid),
        "masukan": {str(p.relative_to(AKAR)): sha(p) for p in masukan if p.exists()},
        "pustaka": {"python": sys.version.split()[0], "pandas": pd.__version__, "numpy": np.__version__,
                    "scipy": scipy.__version__, "scikit-learn": sklearn.__version__,
                    "matplotlib": matplotlib.__version__, "krippendorff": getattr(krippendorff, "__version__", "?")}},
        ensure_ascii=False, indent=2), encoding="utf-8")
    print("\n".join(L))
    print(f"\nSemua keluaran: {out}")


if __name__ == "__main__":
    main()
